"""Owner templates and all test edits use synthetic data and memory buffers."""
import base64
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pytest

from app import engine, owner_workbook as ow


SOURCE = '''# Comentário original sintético
catalog:
  name: exemplo
  description: Catálogo sintético
  tags: {nm_environment: "{{env}}", groupnm_own_stwd: grupo_antigo}
  schemas:
    - name: silver
      tables:
        - name: vendas
          business_name: Vendas
          description: Descrição da tabela
          technical_extension: {campo_preservado: [um, dois]}
          tags: {nm_dom: dominio_teste, fl_mask: "false", nm_class_inf: antiga}
          columns:
            - name: codigo
              type: STRING
              description: Identificador
              tags: {fl_atr_pii: "false", nm_atr_class_inf: interna, personalizada: original}
            - name: observacao
              data_type: STRING
              description: Texto livre
              tags: {fl_atr_mask: true}
'''
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
TAG = "{" + NS["s"] + "}"


def upload(source=SOURCE):
    return {"name": "exemplo.yaml", "content": base64.b64encode(source.encode()).decode()}


def workbook(source=SOURCE):
    files = [upload(source)]
    batch = engine.parse_batch(files)
    assert not batch.errors
    return ow.build_owner_workbook(files, batch.nodes, [])


def inspect(raw):
    return openpyxl.load_workbook(BytesIO(raw), data_only=False)


def header_col(raw, name, header):
    book = inspect(raw)
    try:
        return next(cell.column_letter for cell in book[name][ow.HEADER_ROW + 1] if cell.value == header)
    finally:
        book.close()


def find_row(raw, name, header, expected):
    book = inspect(raw)
    try:
        headers = [c.value for c in book[name][ow.HEADER_ROW + 1]]
        col = headers.index(header)
        return next(row[0].row for row in book[name].iter_rows(min_row=ow.HEADER_ROW + 2)
                    if row[col].value == expected)
    finally:
        book.close()


def edit(raw, name, changes=None, remove_row=None, duplicate_row=None, reorder=False):
    """Edit OOXML directly, since openpyxl.save may use temporary files."""
    with ZipFile(BytesIO(raw)) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    meta = ET.fromstring(entries["xl/workbook.xml"])
    index = next(index + 1 for index, item in enumerate(meta.find("s:sheets", NS))
                 if item.get("name") == name)
    path = f"xl/worksheets/sheet{index}.xml"
    root = ET.fromstring(entries[path])
    data = root.find("s:sheetData", NS)
    for coordinate, value in (changes or {}).items():
        cell = root.find(f".//s:c[@r='{coordinate}']", NS)
        assert cell is not None, coordinate
        for child in list(cell):
            cell.remove(child)
        if value is None:
            cell.attrib.pop("t", None)
        elif isinstance(value, str):
            cell.set("t", "inlineStr")
            ET.SubElement(ET.SubElement(cell, TAG + "is"), TAG + "t").text = value
        else:
            cell.attrib.pop("t", None)
            ET.SubElement(cell, TAG + "v").text = str(value)
    if remove_row:
        data.remove(data.find(f"s:row[@r='{remove_row}']", NS))
    if duplicate_row:
        row = deepcopy(data.find(f"s:row[@r='{duplicate_row}']", NS))
        new_index = max(int(r.get("r")) for r in data) + 1
        row.set("r", str(new_index))
        for cell in row:
            cell.set("r", re_coordinate(cell.get("r"), new_index))
        data.append(row)
        dimension = root.find("s:dimension", NS)
        dimension.set("ref", re_coordinate(dimension.get("ref"), new_index))
    if reorder:
        rows = [row for row in data if int(row.get("r")) > ow.HEADER_ROW + 1]
        for row in rows:
            data.remove(row)
        for index, row in enumerate(reversed(rows), ow.HEADER_ROW + 2):
            row.set("r", str(index))
            for cell in row:
                cell.set("r", re_coordinate(cell.get("r"), index))
            data.append(row)
    entries[path] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for path, content in entries.items():
            archive.writestr(path, content)
    return output.getvalue()


def re_coordinate(coordinate, row):
    import re
    return re.sub(r"\d+$", str(row), coordinate)


def test_exactly_four_friendly_visible_sheets_and_no_approval_gate():
    raw = workbook()
    book = inspect(raw)
    try:
        assert [s.title for s in book if s.sheet_state == "visible"] == ow.VISIBLE_SHEETS
        assert book[ow.PAYLOAD_SHEET].sheet_state == "veryHidden"
        assert book[ow.LIST_SHEET].sheet_state == "veryHidden"
        for name in ow.VISIBLE_SHEETS[1:]:
            ws = book[name]
            assert ws.freeze_panes
            assert ws.auto_filter.ref
            assert ws.protection.sheet
        guide = " ".join(str(cell.value or "") for row in book["Comece aqui"] for cell in row)
        assert ow.DELETE_TOKEN in guide
        assert "branco" in guide
        assert "não é criptografia" in guide
        assert "Sim" in guide and "false" in guide
    finally:
        book.close()


def test_unedited_workbook_recovers_original_sources_without_source_reupload():
    result = ow.read_owner_workbook(workbook())
    assert result["files"] == [upload()]
    assert result["edits"] == []
    assert result["notes"] == []
    assert any(w["code"] == "not_allowed" for w in result["warnings"])
    assert any(w["code"] == "external_list_missing" for w in result["warnings"])
    assert any(w["code"] == "unknown_tag" for w in result["warnings"])
    original = engine.parse_batch(result["files"])
    assert original.nodes[0]["tags"]["nm_environment"] == "{{env}}"
    assert original.nodes[0]["tags"]["groupnm_own_stwd"] == "grupo_antigo"


def test_proposals_start_blank_and_only_existing_tag_cells_are_editable():
    raw = workbook()
    book = inspect(raw)
    try:
        ws = book["Colunas"]
        row = find_row(raw, "Colunas", "Coluna", "codigo")
        headers = [cell.value for cell in ws[ow.HEADER_ROW + 1]]
        for col, header in enumerate(headers, 1):
            if header and "proposto" in header.lower():
                assert ws.cell(row, col).value in ("", None)
        existing = headers.index("Dado pessoal — proposto") + 1
        absent = headers.index("Dado mascarado — proposto") + 1
        assert ws.cell(row, existing).protection.locked is False
        assert ws.cell(row, absent).protection.locked is True
        assert ws.cell(row, existing).fill.fgColor.rgb == "FFFFF2C6"
        # The Owner can now find column tags in Tags as well as in Colunas.
        assert any(row[0].value == "Coluna" for row in book["Tags"].iter_rows(min_row=5))
    finally:
        book.close()


@pytest.mark.parametrize("proposal,expected", [("Sim", "true"), ("Não", "false"), ("true", "true"), ("false", "false")])
def test_boolean_tags_have_owner_choices_and_canonical_yaml(proposal, expected):
    raw = workbook()
    row = find_row(raw, "Colunas", "Coluna", "codigo")
    col = header_col(raw, "Colunas", "Dado pessoal — proposto")
    result = ow.read_owner_workbook(edit(raw, "Colunas", {f"{col}{row}": proposal}))
    changes = [e for e in result["edits"] if e.get("tag") == "fl_atr_pii"]
    assert changes == ([] if expected == "false" else [{"id": engine.parse_batch([upload()]).nodes[3]["id"],
                                                       "field": "tags", "tag": "fl_atr_pii", "value": "true"}])
    book = inspect(raw)
    try:
        assert any(c.value == "Sim" for row in book[ow.LIST_SHEET] for c in row)
        assert any(c.value == "Não" for row in book[ow.LIST_SHEET] for c in row)
    finally:
        book.close()


def test_description_business_name_deletions_notes_and_tag_unknown_roundtrip():
    raw = workbook()
    table_row = find_row(raw, "Objetos", "Nome técnico", "vendas")
    raw = edit(raw, "Objetos", {f"F{table_row}": "Vendas mensais", f"H{table_row}": ow.DELETE_TOKEN,
                               f"I{table_row}": "Conferir a descrição com a área."})
    row = find_row(raw, "Colunas", "Coluna", "codigo")
    col = header_col(raw, "Colunas", "Tag: personalizada — proposto")
    raw = edit(raw, "Colunas", {f"G{row}": "Código comercial", f"{col}{row}": "alterada"})
    result = ow.read_owner_workbook(raw)
    assert len(result["edits"]) == 4
    assert any(e["field"] == "description" and e["value"] is None for e in result["edits"])
    assert result["notes"][0]["comment"] == "Conferir a descrição com a área."
    assert any(w["source"] == "proposal" and w["code"] == "unknown_tag" for w in result["warnings"])
    batch = engine.parse_batch(result["files"])
    reports = engine.apply_edits(batch, result["edits"])
    assert all(r["status"] == "applied" for r in reports)
    with ZipFile(BytesIO(engine.render_yaml_zip(batch, reports))) as archive:
        text = archive.read("exemplo.yaml").decode()
    assert "# Comentário original sintético" in text
    assert "technical_extension" in text
    assert "grupo_antigo" in text
    assert "{{env}}" in text
    assert "Vendas mensais" in text


def test_missing_external_list_is_advisory_for_a_new_existing_tag_value():
    raw = workbook()
    row = find_row(raw, "Tags", "Tag técnica", "nm_dom")
    result = ow.read_owner_workbook(edit(raw, "Tags", {f"F{row}": "novo_dominio"}))
    assert result["edits"][0]["value"] == "novo_dominio"
    assert any(w["source"] == "proposal" and w["code"] == "external_list_missing" for w in result["warnings"])


@pytest.mark.parametrize("kind", ["immutable", "unknown", "missing", "duplicate", "enum", "nontext", "absent_tag"])
def test_changed_identity_bad_proposal_and_row_integrity_fail_the_whole_import(kind):
    raw = workbook()
    row = find_row(raw, "Colunas", "Coluna", "codigo")
    if kind == "immutable":
        changed = edit(raw, "Colunas", {f"F{row}": "Alterou o original"})
    elif kind == "unknown":
        col = header_col(raw, "Colunas", "ID interno")
        changed = edit(raw, "Colunas", {f"{col}{row}": "identificador_forjado"})
    elif kind == "missing":
        changed = edit(raw, "Colunas", remove_row=row)
    elif kind == "duplicate":
        changed = edit(raw, "Colunas", duplicate_row=row)
    elif kind == "enum":
        col = header_col(raw, "Colunas", "Classificação da informação — proposto")
        changed = edit(raw, "Colunas", {f"{col}{row}": "nao_existe"})
    elif kind == "nontext":
        changed = edit(raw, "Colunas", {f"G{row}": 42})
    else:
        col = header_col(raw, "Colunas", "Dado mascarado — proposto")
        changed = edit(raw, "Colunas", {f"{col}{row}": "Sim"})
    with pytest.raises(ValueError):
        ow.read_owner_workbook(changed)


def test_entire_rows_can_be_sorted_without_losing_their_identity():
    raw = workbook()
    row = find_row(raw, "Colunas", "Coluna", "codigo")
    changed = edit(raw, "Colunas", {f"G{row}": "Código revisado"}, reorder=True)
    result = ow.read_owner_workbook(changed)
    assert len(result["edits"]) == 1
    assert result["edits"][0]["value"] == "Código revisado"


def test_extensive_original_description_is_readable_and_preserved_in_payload():
    source = SOURCE.replace("Descrição da tabela", "Descrição longa " * 3000)
    raw = workbook(source)
    row = find_row(raw, "Objetos", "Nome técnico", "vendas")
    book = inspect(raw)
    try:
        assert len(book["Objetos"].cell(row, 7).value) <= 32_000
        assert "preservado" in book["Objetos"].cell(row, 7).value
        assert book["Objetos"].row_dimensions[row].height >= 150
    finally:
        book.close()
    result = ow.read_owner_workbook(raw)
    assert result["edits"] == []
    assert result["files"] == [upload(source)]


def test_native_boolean_is_preserved_when_blank_and_repaired_only_on_proposal():
    raw = workbook()
    row = find_row(raw, "Colunas", "Coluna", "observacao")
    col = header_col(raw, "Colunas", "Dado mascarado — proposto")
    unchanged = ow.read_owner_workbook(raw)
    source_batch = engine.parse_batch(unchanged["files"])
    assert source_batch.nodes[-1]["tags"]["fl_atr_mask"] is True
    result = ow.read_owner_workbook(edit(raw, "Colunas", {f"{col}{row}": "Sim"}))
    assert result["edits"][-1]["value"] == "true"


def test_building_and_reading_need_no_filesystem_or_temporary_files():
    with patch("builtins.open", side_effect=AssertionError("disk")), \
         patch("tempfile.mkstemp", side_effect=AssertionError("temporary")):
        raw = workbook()
        assert ow.read_owner_workbook(raw)["edits"] == []


def legacy_workbook_short_labels():
    with patch.object(ow, "VERSION", ow.LEGACY_VERSION):
        raw = workbook()
    return edit(raw, "Colunas", {"C5": "vendas", "C6": "vendas"})


def test_legacy_short_table_label_preserves_owner_proposals():
    raw = edit(legacy_workbook_short_labels(), "Colunas", {"G5": "Descrição revisada pelo Owner"})
    result = ow.read_owner_workbook(raw)
    assert result["files"] == [upload()]
    assert len(result["edits"]) == 1
    assert result["edits"][0]["value"] == "Descrição revisada pelo Owner"


def test_legacy_qualified_table_label_is_still_accepted():
    with patch.object(ow, "VERSION", ow.LEGACY_VERSION):
        raw = workbook()
    assert ow.read_owner_workbook(raw)["edits"] == []


@pytest.mark.parametrize("header,value", [("Tabela / visão", "outra_tabela"),
                                         ("SHA256 de origem", "hash_invalido"),
                                         ("Coluna", "outra_coluna")])
def test_legacy_compatibility_does_not_allow_identity_tampering(header, value):
    raw = legacy_workbook_short_labels()
    col = header_col(raw, "Colunas", header)
    with pytest.raises(ValueError, match="alterado"):
        ow.read_owner_workbook(edit(raw, "Colunas", {f"{col}5": value}))


def test_v2_requires_the_qualified_table_label():
    raw = edit(workbook(), "Colunas", {"C5": "vendas"})
    with pytest.raises(ValueError, match="alterado"):
        ow.read_owner_workbook(raw)
