"""The two-step owner hand-off works from workbook bytes alone, in RAM."""
from __future__ import annotations

import base64
from contextlib import ExitStack
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile
import zlib

import openpyxl
import pytest

from app import service
from app import engine
from app.engine import parse_batch


SOURCE = '''# Contrato sintético preservado
table:
  id: catalogo_teste.silver.pedidos
  name: pedidos
  description: Descrição inicial
  business_name: Pedidos
  format: delta  # formato técnico preservado
  custom: {preservar: sim}
  tags:
    fl_mask: "false"
    fl_apt_cons: "false"
    fl_metadata_ai_generate: "true"
  columns:
    - name: codigo
      type: STRING
      nullable: false
      description: Código original
      tags: {fl_atr_mask: "false", fl_atr_pii: "false"}
ingest:
  source: {kind: synthetic, options: {marker: preservar}}
'''


def upload(raw=SOURCE, name="pedidos.yaml"):
    if isinstance(raw, str):
        raw = raw.encode("utf-8")
    return {"name": name, "content": base64.b64encode(raw).decode("ascii")}


def make_zip(entries):
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, raw in entries:
            archive.writestr(name, raw)
    return output.getvalue()


def convert(content, name="revisao.xlsx"):
    return service.convert_owner_excels([{"name": name, "content": content}])


def output_files(result):
    with ZipFile(BytesIO(result["content"])) as archive:
        assert archive.namelist()
        assert all(name.lower().endswith((".yaml", ".yml")) for name in archive.namelist())
        return {name: archive.read(name) for name in archive.namelist()}


def workbook_edit(content, sheet, row, column, value, *, formula=False):
    """Patch OOXML in memory; openpyxl.save may create temporary files."""
    from xml.etree import ElementTree as ET

    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    tag = "{" + ns["s"] + "}"
    with ZipFile(BytesIO(content)) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    book = ET.fromstring(entries["xl/workbook.xml"])
    index = next(i + 1 for i, item in enumerate(book.find("s:sheets", ns)) if item.get("name") == sheet)
    path = f"xl/worksheets/sheet{index}.xml"
    root = ET.fromstring(entries[path])
    coordinate = openpyxl.utils.get_column_letter(column) + str(row)
    cell = root.find(f".//s:c[@r='{coordinate}']", ns)
    if cell is None:
        row_element = root.find(f".//s:row[@r='{row}']", ns)
        assert row_element is not None
        cell = ET.SubElement(row_element, tag + "c", {"r": coordinate})
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)
    if formula:
        ET.SubElement(cell, tag + "f").text = value
    elif value is not None:
        cell.set("t", "inlineStr")
        ET.SubElement(ET.SubElement(cell, tag + "is"), tag + "t").text = value
    entries[path] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    return make_zip(entries.items())


def locate_edit(content, *, sheet, name, field="description", tag=None):
    """Find a real owner-facing field, keeping tests independent of row order."""
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    try:
        ws = workbook[sheet]
        target_words = ("descrição", "descriçao") if field == "description" else ("valor",)
        for row in ws.iter_rows(min_row=1, max_row=min(10, ws.max_row)):
            headers = [str(cell.value or "").lower() for cell in row]
            candidates = [i for i, text in enumerate(headers) if any(word in text for word in target_words)]
            candidates = [i for i in candidates if any(word in headers[i] for word in ("propost", "revis", "valid"))] or candidates
            if not candidates:
                continue
            editable = candidates[-1]
            for data in ws.iter_rows(min_row=row[0].row + 1):
                values = [str(cell.value or "") for cell in data]
                if not any(value == name or value.endswith("." + name) for value in values):
                    continue
                if tag is not None and tag not in values:
                    continue
                if data[editable].protection.locked:
                    continue
                return data[editable].row, editable + 1
        raise AssertionError(f"Editable {field} row not found in {sheet}")
    finally:
        workbook.close()


def edit_owner_field(content, *, sheet, name, value, field="description", tag=None):
    row, column = locate_edit(content, sheet=sheet, name=name, field=field, tag=tag)
    return workbook_edit(content, sheet, row, column, value)


def assert_same_documents(original, updated):
    before = parse_batch(original)
    after = parse_batch([upload(value, name) for name, value in updated.items()])
    assert not before.errors and not after.errors
    assert {d.filename: d.roots for d in before.documents} == {d.filename: d.roots for d in after.documents}


def test_owner_excel_is_self_contained_and_roundtrips_every_file_without_approval():
    files = [upload()]
    excel = service.create_owner_excel(files)
    result = convert(excel["content"])
    updated = output_files(result)
    assert set(updated) == {"pedidos.yaml"}
    assert_same_documents(files, updated)
    text = updated["pedidos.yaml"].decode("utf-8")
    assert "# Contrato sintético preservado" in text
    assert "# formato técnico preservado" in text
    assert "marker: preservar" in text
    assert result["changes"] == []


def test_batch_zip_roundtrips_nested_paths_and_yaml_yml_extensions():
    entries = [("grupo/pedidos.yaml", SOURCE), ("outro.yml", SOURCE.replace("pedidos", "clientes"))]
    files = [upload(make_zip(entries), "lote.zip")]
    excel = service.create_owner_excel(files)
    updated = output_files(convert(excel["content"]))
    assert set(updated) == {"lote/grupo/pedidos.yaml", "lote/outro.yml"}
    assert_same_documents(files, updated)


def test_existing_governance_warnings_do_not_block_unchanged_export():
    source = SOURCE.replace('fl_mask: "false"', 'fl_mask: "valor_antigo"').replace('fl_metadata_ai_generate: "true"', 'nm_dom: "dominio_sintetico"')
    files = [upload(source)]
    excel = service.create_owner_excel(files)
    updated = output_files(convert(excel["content"]))
    assert_same_documents(files, updated)


def test_owner_proposals_apply_descriptions_tags_and_preserve_other_fields():
    excel = service.create_owner_excel([upload()])["content"]
    excel = edit_owner_field(excel, sheet="Objetos", name="pedidos", value="Descrição revisada")
    excel = edit_owner_field(excel, sheet="Colunas", name="codigo", value="Código revisado")
    excel = edit_owner_field(excel, sheet="Tags", name="catalogo_teste.silver.pedidos", field="tags", tag="fl_metadata_ai_generate", value="false")
    result = convert(excel)
    updated = output_files(result)
    parsed = parse_batch([upload(updated["pedidos.yaml"])])
    obj = next(node for node in parsed.nodes if node["level"] == "Table")
    col = next(node for node in parsed.nodes if node["level"] == "Column")
    assert obj["description"] == "Descrição revisada"
    assert col["description"] == "Código revisado"
    assert obj["tags"]["fl_metadata_ai_generate"] == "false"
    assert obj["business_name"] == "Pedidos" and col["type"] == "STRING"
    assert len(result["changes"]) == 3
    assert "# formato técnico preservado" in updated["pedidos.yaml"].decode()


def test_invalid_boolean_proposal_and_new_mask_inconsistency_reject_whole_export():
    excel = service.create_owner_excel([upload()])["content"]
    invalid = edit_owner_field(excel, sheet="Tags", name="catalogo_teste.silver.pedidos", field="tags", tag="fl_metadata_ai_generate", value="CONTEUDO_PRIVADO_SENTINELA")
    with pytest.raises(ValueError) as exc:
        convert(invalid)
    assert "CONTEUDO_PRIVADO_SENTINELA" not in str(exc.value)
    # A masked column requires the table flag. The reverse is not a documented rule.
    workbook = openpyxl.load_workbook(BytesIO(excel), read_only=True)
    headers = list(workbook["Colunas"].iter_rows(min_row=4, max_row=4, values_only=True))[0]
    mask_column = headers.index("Dado mascarado — proposto") + 1
    workbook.close()
    inconsistent = workbook_edit(excel, "Colunas", 5, mask_column, "Sim")
    with pytest.raises(ValueError, match="inconsistência"):
        convert(inconsistent)


def test_formulas_and_changes_to_original_cells_are_rejected():
    excel = service.create_owner_excel([upload()])["content"]
    row, column = locate_edit(excel, sheet="Objetos", name="pedidos")
    formula = workbook_edit(excel, "Objetos", row, column, 'HYPERLINK("https://example.invalid","synthetic")', formula=True)
    with pytest.raises(ValueError, match="[Ff]órmula") as caught:
        convert(formula)
    coordinate = openpyxl.utils.get_column_letter(column) + str(row)
    assert f'Aba "Objetos", célula {coordinate}' in str(caught.value)
    assert "Colar especial > Valores" in str(caught.value)
    assert "HYPERLINK" not in str(caught.value)
    changed_original = workbook_edit(excel, "Objetos", row, column - 1, "Alteração no campo protegido")
    with pytest.raises(ValueError, match="alterado"):
        convert(changed_original)


def test_literal_text_starting_with_equals_is_not_a_formula():
    excel = service.create_owner_excel([upload()])["content"]
    row, column = locate_edit(excel, sheet="Objetos", name="pedidos")
    revised = workbook_edit(excel, "Objetos", row, column, "=texto literal, não calculado")
    result = convert(revised)
    parsed = parse_batch([upload(raw, name) for name, raw in output_files(result).items()])
    assert parsed.nodes[0]["description"] == "=texto literal, não calculado"


def test_owner_notes_are_visible_to_governance_and_never_injected_into_yaml():
    excel = service.create_owner_excel([upload()])["content"]
    row, _ = locate_edit(excel, sheet="Objetos", name="pedidos")
    excel = workbook_edit(excel, "Objetos", row, 9, "Observação sintética para governança")
    result = convert(excel)
    assert any("Observação sintética para governança" in str(warning) for warning in result["warnings"])
    assert "Observação sintética para governança" not in output_files(result)["pedidos.yaml"].decode()


def test_returned_excel_batches_include_all_objects_and_reject_duplicate_versions():
    first = service.create_owner_excel([upload()])["content"]
    second = service.create_owner_excel([upload(SOURCE.replace("pedidos", "clientes"), "clientes.yml")])["content"]
    result = service.convert_owner_excels([
        {"name": "owner_um.xlsx", "content": first},
        {"name": "owner_dois.xlsx", "content": second},
    ])
    updated = output_files(result)
    assert len(updated) == 2
    assert {Path(name).name for name in updated} == {"pedidos.yaml", "clientes.yml"}
    with pytest.raises(ValueError, match="mesmo objeto"):
        service.convert_owner_excels([{"name": "um.xlsx", "content": first}, {"name": "dois.xlsx", "content": first}])


def test_several_small_workbooks_cannot_bypass_cumulative_restored_source_limit(monkeypatch):
    first = service.create_owner_excel([upload()])["content"]
    second = service.create_owner_excel([upload(SOURCE.replace("pedidos", "clientes"), "clientes.yml")])["content"]
    monkeypatch.setattr(engine, "MAX_BATCH_BYTES", len(SOURCE.encode("utf-8")) + 30)
    with pytest.raises(ValueError):
        service.convert_owner_excels([{"name": "um.xlsx", "content": first}, {"name": "dois.xlsx", "content": second}])


def test_conversion_rejects_partial_bad_source_batches_with_safe_message():
    with pytest.raises(ValueError) as exc:
        service.create_owner_excel([upload(), upload("table: [CONTEUDO_PRIVADO_SENTINELA", "invalido.yaml")])
    assert "CONTEUDO_PRIVADO_SENTINELA" not in str(exc.value)


@pytest.mark.parametrize("raw", [b"not-a-zip", b"PK\x03\x04broken", b""])
def test_invalid_excel_is_a_safe_user_error(raw):
    with pytest.raises(ValueError) as exc:
        convert(raw)
    assert "Traceback" not in str(exc.value)


def test_corrupt_deflate_yaml_zip_is_a_safe_user_error():
    packed = bytearray(make_zip([("pedidos.yaml", SOURCE)]))
    with ZipFile(BytesIO(packed)) as archive:
        info = archive.infolist()[0]
        offset = info.header_offset + 30 + len(info.filename.encode()) + len(info.extra)
    packed[offset] = 0x06  # DEFLATE reserved block type, never decompressed successfully.
    with pytest.raises(ValueError):
        service.create_owner_excel([upload(bytes(packed), "corrompido.zip")])


def test_complete_new_journey_has_no_disk_or_temporary_file_access():
    packed = upload(make_zip([("grupo/pedidos.yaml", SOURCE)]), "lote.zip")
    with ExitStack() as stack:
        for target in ("builtins.open", "tempfile.NamedTemporaryFile", "tempfile.TemporaryFile", "tempfile.mkstemp", "tempfile.SpooledTemporaryFile"):
            stack.enter_context(patch(target, side_effect=AssertionError("Persistent I/O attempted")))
        for method in ("write_bytes", "write_text"):
            stack.enter_context(patch.object(Path, method, side_effect=AssertionError("Persistent write attempted")))
        excel = service.create_owner_excel([packed])
        updated = output_files(convert(excel["content"]))
        assert "lote/grupo/pedidos.yaml" in updated


def test_xlsx_macros_external_links_and_entities_are_rejected():
    excel = service.create_owner_excel([upload()])["content"]
    with ZipFile(BytesIO(excel)) as archive:
        entries = [(name, archive.read(name)) for name in archive.namelist()]
    malicious_entries = [
        ("xl/vbaProject.bin", b"synthetic macro marker"),
        ("xl/externalLinks/externalLink1.xml", b"<externalLink/>"),
        ("evil.xml", b'<!DOCTYPE x [<!ENTITY private "CONTEUDO_PRIVADO_SENTINELA">]><x>&private;</x>'),
        ("../evil.xml", b"<x/>"),
    ]
    for addition in malicious_entries:
        with pytest.raises(ValueError) as exc:
            convert(make_zip(entries + [addition]))
        assert "CONTEUDO_PRIVADO_SENTINELA" not in str(exc.value)


def test_xlsx_high_expansion_is_rejected_before_decompressing():
    packed = make_zip([("xl/worksheets/sheet1.xml", b"x" * (2 * 1024 * 1024))])
    with patch.object(ZipFile, "read", side_effect=AssertionError("Decompression before preflight")):
        with pytest.raises(ValueError):
            convert(packed)


def payload_parts(content):
    book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
    try:
        ws = book["_metadata"]
        manifest = json.loads(ws["A1"].value)
        packed = base64.b64decode("".join(ws.cell(row, 1).value for row in range(2, manifest["chunks"] + 2)))
        return manifest, packed
    finally:
        book.close()


@pytest.mark.parametrize("mode", ["base64", "checksum", "truncated", "trailing", "bomb", "unsafe_source"])
def test_portable_payload_corruption_or_expansion_is_rejected_safely(mode):
    excel = service.create_owner_excel([upload()])["content"]
    manifest, packed = payload_parts(excel)
    assert manifest["chunks"] == 1
    chunk = None
    if mode == "base64":
        chunk = "not base64: CONTEUDO_PRIVADO_SENTINELA"
    elif mode == "checksum":
        manifest["sha256"] = "0" * 64
    elif mode == "truncated":
        chunk = base64.b64encode(packed[:-1]).decode()
    elif mode == "trailing":
        chunk = base64.b64encode(packed + zlib.compress(b"trailing private marker")).decode()
    elif mode == "bomb":
        manifest["length"] = 20
        chunk = base64.b64encode(zlib.compress(b"x" * (2 * 1024 * 1024))).decode()
    elif mode == "unsafe_source":
        payload = json.loads(zlib.decompress(packed))
        payload["files"] = [upload("table: [CONTEUDO_PRIVADO_SENTINELA")]
        raw = json.dumps(payload).encode()
        manifest.update(length=len(raw), sha256=sha256(raw).hexdigest())
        chunk = base64.b64encode(zlib.compress(raw)).decode()
    excel = workbook_edit(excel, "_metadata", 1, 1, json.dumps(manifest))
    if chunk is not None:
        excel = workbook_edit(excel, "_metadata", 2, 1, chunk)
    with pytest.raises(ValueError) as exc:
        convert(excel)
    assert "CONTEUDO_PRIVADO_SENTINELA" not in str(exc.value)
