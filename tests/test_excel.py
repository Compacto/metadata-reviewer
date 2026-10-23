"""Synthetic fixtures only. Workbook editing helpers also operate in RAM."""
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
import pytest

from app import rules
from app.excel import DELETE_TOKEN, HEADER_ROW, STATUSES, VISIBLE_SHEETS, export_workbook, preview_workbook

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
TAG = "{" + NS["s"] + "}"
REFS = {"nm_dom": ["dominio_teste"], "nm_subdom": ["subdominio_teste"], "nm_sist_orig": ["sistema_teste"], "ct_area_neg": ["area_teste"]}


@pytest.fixture
def nodes():
    def node(node_id, level, path, tags, description="Descrição sintética"):
        parts = path.split(".")
        return {"id": node_id, "level": level, "path": path, "catalog": parts[0],
                "schema": parts[1] if len(parts) > 1 else "", "name": parts[-1],
                "object_id": "table-1" if level == "Column" else node_id,
                "filename": "exemplo_sintetico.yaml", "source_sha256": "a" * 64,
                "description": description, "business_name": "Nome sintético" if level == "Table" else None,
                "tags": tags, "original_tags": deepcopy(tags), "original_description": description,
                "original_business_name": "Nome sintético" if level == "Table" else None, "data_type": "STRING" if level == "Column" else None}
    return [node("catalog-1", "Catalog", "catalogo_teste", {"nm_class_inf": "restrita"}),
            node("schema-1", "Schema", "catalogo_teste.silver", {"nm_class_inf": "restrita"}),
            node("table-1", "Table", "catalogo_teste.silver.tabela_teste", {"fl_mask": "false", "nm_class_inf": "interna", "nm_dom": "dominio_teste", "fl_apt_cons": "false"}),
            node("view-1", "View", "catalogo_teste.silver.view_teste", {}),
            node("column-1", "Column", "catalogo_teste.silver.tabela_teste.coluna_teste", {"fl_atr_mask": "false", "nm_atr_class_inf": "interna", "fl_atr_pii": "false"})]


def exported(nodes, reviews=None, references=REFS):
    return export_workbook(nodes, reviews or {}, rules.validate_nodes(nodes, references), references)


def reader(content):
    return openpyxl.load_workbook(BytesIO(content), data_only=False)


def row_number(content, sheet, node_id, tag=None):
    book = reader(content)
    try:
        ws = book[sheet]
        headers = [cell.value for cell in ws[HEADER_ROW + 1]]
        for row in ws.iter_rows(min_row=HEADER_ROW + 2):
            if row[headers.index("ID estável")].value == node_id:
                if tag is None or row[headers.index("Tag técnica")].value == tag:
                    return row[0].row
        raise AssertionError("Row not found")
    finally:
        book.close()


def column_letter(content, sheet, header):
    book = reader(content)
    try:
        return next(cell.column_letter for cell in book[sheet][HEADER_ROW + 1] if cell.value == header)
    finally:
        book.close()


def mutate(content, sheet, edits=None, *, delete_row=None, duplicate_row=None, formula=None, numeric=None):
    """Patch OOXML bytes directly; never call openpyxl.save (which uses temp files)."""
    with ZipFile(BytesIO(content)) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    workbook = ET.fromstring(entries["xl/workbook.xml"])
    sheets = workbook.find("s:sheets", NS)
    index = next(i + 1 for i, entry in enumerate(sheets) if entry.get("name") == sheet)
    path = f"xl/worksheets/sheet{index}.xml"
    root = ET.fromstring(entries[path])
    data = root.find("s:sheetData", NS)
    for cell_id, value in (edits or {}).items():
        cell = root.find(f".//s:c[@r='{cell_id}']", NS)
        assert cell is not None
        for child in list(cell):
            cell.remove(child)
        if value is None:
            cell.attrib.pop("t", None)
        else:
            cell.set("t", "inlineStr")
            ET.SubElement(ET.SubElement(cell, TAG + "is"), TAG + "t").text = value
    if formula:
        cell = root.find(f".//s:c[@r='{formula}']", NS)
        for child in list(cell):
            cell.remove(child)
        cell.attrib.pop("t", None)
        ET.SubElement(cell, TAG + "f").text = "1+1"
    if numeric:
        cell = root.find(f".//s:c[@r='{numeric}']", NS)
        for child in list(cell):
            cell.remove(child)
        cell.set("t", "b")
        ET.SubElement(cell, TAG + "v").text = "1"
    if delete_row:
        data.remove(data.find(f"s:row[@r='{delete_row}']", NS))
    if duplicate_row:
        row = deepcopy(data.find(f"s:row[@r='{duplicate_row}']", NS))
        new_number = max(int(item.get("r")) for item in data) + 1
        row.set("r", str(new_number))
        for cell in row:
            cell.set("r", "".join(c for c in cell.get("r") if c.isalpha()) + str(new_number))
        data.append(row)
        dimension = root.find("s:dimension", NS)
        last_col = "".join(c for c in dimension.get("ref").split(":")[-1] if c.isalpha())
        dimension.set("ref", "A1:" + last_col + str(new_number))
    entries[path] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)
    return output.getvalue()


def edit_field(content, sheet, node_id, header, value, tag=None):
    target = column_letter(content, sheet, header) + str(row_number(content, sheet, node_id, tag))
    return mutate(content, sheet, {target: value})


def test_export_five_visible_sheets_native_validations_and_protection(nodes):
    with patch("tempfile.mkstemp", side_effect=AssertionError("No temporary files")), patch("tempfile.NamedTemporaryFile", side_effect=AssertionError("No temporary files")):
        content = exported(nodes)
        result = preview_workbook(content, nodes, {}, REFS)
    assert not result["edits"] and not result["rejected"] and not result["missing"]
    book = reader(content)
    assert [s.title for s in book if s.sheet_state == "visible"] == VISIBLE_SHEETS
    assert book["_Listas"].sheet_state == "hidden"
    assert book.defined_names
    for sheet in (book["Objects"], book["Tags"], book["Columns"], book["Approval"]):
        assert sheet.freeze_panes and sheet.auto_filter.ref and sheet.protection.sheet
        assert sheet.data_validations.count
        assert all(validation.formula1.removeprefix("=") in book.defined_names for validation in sheet.data_validations.dataValidation)
    assert book["Objects"]["F5"].protection.locked is False
    assert book["Objects"]["A5"].protection.locked is True
    book.close()


def test_all_levels_absent_values_and_missing_lists_are_explicit(nodes):
    content = exported(nodes, references=None)
    book = reader(content)
    assert {row[1].value for row in book["Tags"].iter_rows(min_row=5)} == {"Catalog", "Schema", "Table", "View", "Column"}
    row = row_number(content, "Tags", "column-1", "fl_atr_sens")
    assert book["Tags"].cell(row, 6).value == "Não"
    assert book["Tags"].cell(row, 7).value in (None, "")
    assert book["Tags"].cell(row, 8).value in (None, "")
    all_text = " ".join(str(cell.value or "") for row in book["Read Me"] for cell in row)
    assert "Lista oficial não carregada" in all_text
    assert "Não identificada" in all_text and "Information Schema" in all_text
    assert "ct_risco*" in all_text
    book.close()


def test_description_business_name_and_explicit_tag_clear(nodes):
    content = exported(nodes)
    content = edit_field(content, "Objects", "table-1", "Descrição proposta", "Descrição revisada")
    content = edit_field(content, "Objects", "table-1", "Nome de negócio proposto", "Novo nome")
    content = edit_field(content, "Tags", "column-1", "Valor proposto", DELETE_TOKEN, "fl_atr_pii")
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["rejected"]
    assert {("table-1", "description", "Descrição revisada"), ("table-1", "business_name", "Novo nome"), ("column-1", "tags", None)} == {(e["id"], e["field"], e["value"]) for e in result["edits"]}
    assert nodes[-1]["tags"]["fl_atr_pii"] == "false"  # preview never mutates


def test_conflicting_column_and_tags_edits_are_rejected(nodes):
    content = exported(nodes)
    # Both differ from baseline false: clearing vs setting true.
    content = edit_field(content, "Tags", "column-1", "Valor proposto", DELETE_TOKEN, "fl_atr_pii")
    content = edit_field(content, "Columns", "column-1", "fl_atr_pii proposto", "true")
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["edits"]
    assert any(row["code"] == "conflict" for row in result["rejected"])


def test_same_proposal_twice_is_accepted_once(nodes):
    content = exported(nodes)
    content = edit_field(content, "Tags", "column-1", "Valor proposto", "true", "fl_atr_pii")
    content = edit_field(content, "Columns", "column-1", "fl_atr_pii proposto", "true")
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["rejected"] and len(result["edits"]) == 1


@pytest.mark.parametrize("value", ["yes", "TRUE", "False", "verdadeiro"])
def test_invalid_controlled_values_rejected(nodes, value):
    content = edit_field(exported(nodes), "Tags", "column-1", "Valor proposto", value, "fl_atr_pii")
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["edits"] and result["rejected"][0]["code"] == "invalid_proposal"


def test_native_boolean_cell_is_not_a_valid_tag_string(nodes):
    content = exported(nodes)
    target = column_letter(content, "Tags", "Valor proposto") + str(row_number(content, "Tags", "column-1", "fl_atr_pii"))
    result = preview_workbook(mutate(content, "Tags", numeric=target), nodes, {}, REFS)
    assert any(item["code"] == "value_type" for item in result["rejected"])
    assert not result["edits"]


def test_old_source_fingerprint_and_immutable_edits_rejected(nodes):
    content = edit_field(exported(nodes), "Objects", "table-1", "Descrição proposta", "Mudança")
    new_nodes = deepcopy(nodes)
    for node in new_nodes:
        node["source_sha256"] = "b" * 64
    result = preview_workbook(content, new_nodes, {}, REFS)
    assert not result["edits"] and any(item["code"] == "immutable" for item in result["rejected"])
    content = edit_field(content, "Objects", "table-1", "Nome técnico", "adulterado")
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["edits"] and any(item["code"] == "immutable" for item in result["rejected"])


@pytest.mark.parametrize("action", ["missing", "duplicate"])
def test_missing_or_duplicate_tag_rows_block_corresponding_column_edit(nodes, action):
    content = edit_field(exported(nodes), "Columns", "column-1", "fl_atr_pii proposto", "true")
    row = row_number(content, "Tags", "column-1", "fl_atr_pii")
    content = mutate(content, "Tags", **{"delete_row" if action == "missing" else "duplicate_row": row})
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["edits"]
    assert result["missing"] if action == "missing" else any(item["code"] == "duplicate" for item in result["rejected"])


def test_formula_injection_rejected_and_export_treats_formula_like_text(nodes):
    changed = deepcopy(nodes)
    changed[2]["description"] = '=HYPERLINK("https://invalid.example","x")'
    content = exported(changed)
    book = reader(content)
    assert book["Objects"]["H5"].data_type == "s"
    book.close()
    result = preview_workbook(mutate(content, "Objects", formula="H5"), changed, {}, REFS)
    assert not result["edits"] and result["rejected"][0]["code"] == "formula"


def test_new_mask_inconsistency_rejects_tags_but_joint_fix_is_valid(nodes):
    content = edit_field(exported(nodes), "Columns", "column-1", "fl_atr_mask proposto", "true")
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["edits"] and any(item["code"] == "new_consistency_error" for item in result["rejected"])
    content = edit_field(content, "Tags", "table-1", "Valor proposto", "true", "fl_mask")
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["rejected"] and len(result["edits"]) == 2


def test_consumption_approval_rule_rejects_incomplete_proposal(nodes):
    content = edit_field(exported(nodes), "Tags", "table-1", "Valor proposto", "true", "fl_apt_cons")
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["edits"] and any(item["code"] == "new_consistency_error" for item in result["rejected"])


def test_functional_completion_keeps_technical_reservations_and_notes(nodes):
    content = exported(nodes, references=None)
    content = edit_field(content, "Approval", "table-1", "Validation completed", "Sim")
    content = edit_field(content, "Approval", "table-1", "Data Owner", "Owner declarado")
    content = edit_field(content, "Approval", "table-1", "Comentários finais", "Revisado com pendência de lista")
    content = edit_field(content, "Tags", "column-1", "Comentários", "Dúvida do Owner", "fl_atr_pii")
    result = preview_workbook(content, nodes, {}, None)
    assert not result["rejected"]
    assert result["reviews"]["table-1"]["completed"] is True
    assert result["reviews"]["table-1"]["status"] == "Validated with remarks"
    assert result["reviews"]["table-1"]["owner"] == "Owner declarado"
    assert result["reviews"]["table-1"]["final_comments"] == "Revisado com pendência de lista"
    assert result["reviews"]["column-1"]["tag_comments"]["fl_atr_pii"] == "Dúvida do Owner"


def test_owner_can_reopen_previously_completed_object(nodes):
    reviews = {"table-1": {"completed": True, "status": "Validated"}}
    content = edit_field(exported(nodes, reviews), "Approval", "table-1", "Validation completed", "Não")
    result = preview_workbook(content, nodes, reviews, REFS)
    assert result["reviews"]["table-1"]["completed"] is False
    assert result["reviews"]["table-1"]["status"] == "In review"


def test_invalid_or_oversized_container_has_safe_diagnostic(nodes):
    for content in (b"not an excel", b"x" * (20 * 1024 * 1024 + 1)):
        result = preview_workbook(content, nodes, {}, REFS)
        assert not result["edits"] and result["rejected"][0]["code"] == "invalid_workbook"


def test_xml_entity_container_is_rejected_without_parsing(nodes):
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("xl/workbook.xml", '<!DOCTYPE foo [<!ENTITY x SYSTEM "file:///etc/passwd">]><foo/>')
    result = preview_workbook(output.getvalue(), nodes, {}, REFS)
    assert result["rejected"][0]["code"] == "invalid_workbook"
    assert "XML" in result["rejected"][0]["message"]


def test_external_list_change_requires_official_reference(nodes):
    content = exported(nodes, references=None)
    content = edit_field(content, "Tags", "table-1", "Valor proposto", "dominio_inventado", "nm_dom")
    result = preview_workbook(content, nodes, {}, None)
    assert not result["edits"] and "Lista oficial não carregada" in result["rejected"][0]["message"]


def test_blank_cells_preserve_all_existing_business_values_and_review_notes(nodes):
    reviews = {"table-1": {"owner": "Owner sintético", "comments": "Manter comentário", "status": "In review"},
               "column-1": {"tag_comments": {"fl_atr_pii": "Manter nota"}}}
    content = exported(nodes, reviews)
    for sheet, node_id, header, tag in [
        ("Objects", "table-1", "Descrição proposta", None),
        ("Objects", "table-1", "Nome de negócio proposto", None),
        ("Objects", "table-1", "Comentários do Data Owner", None),
        ("Approval", "table-1", "Data Owner", None),
        ("Tags", "column-1", "Valor proposto", "fl_atr_pii"),
        ("Tags", "column-1", "Comentários", "fl_atr_pii"),
        ("Columns", "column-1", "fl_atr_mask proposto", None),
    ]:
        content = edit_field(content, sheet, node_id, header, None, tag)
    result = preview_workbook(content, nodes, reviews, REFS)
    assert not result["edits"] and not result["changes"] and not result["rejected"]
    assert result["reviews"] == reviews


def test_explicit_token_removes_only_existing_editable_fields(nodes):
    reviews = {"table-1": {"comments": "Comentário existente"}}
    content = exported(nodes, reviews)
    content = edit_field(content, "Objects", "table-1", "Descrição proposta", DELETE_TOKEN)
    content = edit_field(content, "Objects", "table-1", "Nome de negócio proposto", DELETE_TOKEN)
    content = edit_field(content, "Objects", "table-1", "Comentários do Data Owner", DELETE_TOKEN)
    result = preview_workbook(content, nodes, reviews, REFS)
    assert not result["rejected"]
    assert all(edit["value"] is None for edit in result["edits"])
    assert {edit["field"] for edit in result["edits"]} == {"description", "business_name"}
    assert result["reviews"]["table-1"]["comments"] == ""


def test_explicit_deletion_of_absent_field_or_status_is_rejected(nodes):
    content = edit_field(exported(nodes), "Tags", "column-1", "Valor proposto", DELETE_TOKEN, "fl_atr_sens")
    content = edit_field(content, "Objects", "table-1", "Status da validação", DELETE_TOKEN)
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["edits"]
    assert {item["code"] for item in result["rejected"]} == {"deletion_absent", "status"}


def test_explicit_deletion_that_breaks_mask_rule_is_rejected(nodes):
    nodes[2]["tags"]["fl_mask"] = nodes[2]["original_tags"]["fl_mask"] = "true"
    nodes[-1]["tags"]["fl_atr_mask"] = nodes[-1]["original_tags"]["fl_atr_mask"] = "true"
    content = edit_field(exported(nodes), "Tags", "table-1", "Valor proposto", DELETE_TOKEN, "fl_mask")
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["edits"] and any(item["code"] == "new_consistency_error" for item in result["rejected"])


@pytest.mark.parametrize("completed", [False, True])
def test_rejected_decision_is_preserved_with_or_without_completion(nodes, completed):
    assert STATUSES == ["Pending", "In review", "Awaiting clarification", "Validated", "Validated with remarks", "Rejected"]
    reviews = {"table-1": {"completed": completed, "status": "In review"}}
    content = edit_field(exported(nodes, reviews), "Approval", "table-1", "Status final", "Rejected")
    result = preview_workbook(content, nodes, reviews, REFS)
    assert not result["rejected"]
    assert result["reviews"]["table-1"]["status"] == "Rejected"
    assert result["reviews"]["table-1"]["completed"] is completed


def test_invalid_duplicate_proposal_blocks_alternative_edit_in_other_tab(nodes):
    content = edit_field(exported(nodes), "Tags", "column-1", "Valor proposto", "INVALID", "fl_atr_pii")
    content = edit_field(content, "Columns", "column-1", "fl_atr_pii proposto", "true")
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["edits"]
    assert {item["code"] for item in result["rejected"]} >= {"invalid_proposal", "blocked_record"}


def test_native_yaml_boolean_requires_explicit_text_correction(nodes):
    nodes[-1]["tags"]["fl_atr_pii"] = nodes[-1]["original_tags"]["fl_atr_pii"] = False
    content = exported(nodes)
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["edits"] and not result["rejected"]
    content = edit_field(content, "Tags", "column-1", "Valor proposto", "false", "fl_atr_pii")
    result = preview_workbook(content, nodes, {}, REFS)
    assert not result["rejected"]
    assert result["edits"] == [{"id": "column-1", "field": "tags", "tag": "fl_atr_pii", "value": "false"}]


def test_awaiting_decision_survives_completed_review_and_owner_pending_means_remarks(nodes):
    content = exported(nodes)
    content = edit_field(content, "Approval", "table-1", "Status final", "Awaiting clarification")
    content = edit_field(content, "Approval", "table-1", "Validation completed", "Sim")
    result = preview_workbook(content, nodes, {}, REFS)
    assert result["reviews"]["table-1"]["status"] == "Awaiting clarification"
    content = exported(nodes)
    content = edit_field(content, "Approval", "table-1", "Validation completed", "Sim")
    content = edit_field(content, "Approval", "table-1", "Pendências do Owner", "Confirmar com área responsável")
    result = preview_workbook(content, nodes, {}, REFS)
    assert result["reviews"]["table-1"]["status"] == "Validated with remarks"
