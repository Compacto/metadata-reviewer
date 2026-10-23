"""Excel review files. All input, authoring, ZIP handling and parsing stay in RAM.

Sheet protection helps owners avoid accidental edits; the importer independently
checks identities and original snapshots, so protection is not a trust boundary.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from copy import deepcopy
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
import json
from pathlib import PurePosixPath
import re
from typing import Any
from zipfile import BadZipFile, ZipFile

import openpyxl
import xlsxwriter
from xlsxwriter.utility import xl_col_to_name

from . import rules


STATUSES = ["Pending", "In review", "Awaiting clarification", "Validated", "Validated with remarks", "Rejected"]
VISIBLE_SHEETS = ["Read Me", "Objects", "Tags", "Columns", "Approval"]
DELETE_TOKEN = "__DELETE__"
COLUMN_TAGS = ["fl_reg_apl", "fl_atr_pii", "nm_atr_pii", "fl_atr_sens", "fl_atr_mask", "nm_atr_class_inf", "fl_metadata_ai_generate"]
HEADER_ROW = 3
MAX_XLSX_BYTES = 20 * 1024 * 1024
MAX_XML_BYTES = 80 * 1024 * 1024
MAX_ROWS = 100_000
MISSING_LIST = "Lista oficial não carregada"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _original(node: dict, field: str) -> Any:
    return node.get("original_" + field, node.get(field))


def _source_hashes(nodes: list[dict]) -> dict[str, str]:
    grouped: dict[str, list] = defaultdict(list)
    result = {}
    for node in nodes:
        filename = node.get("filename", "")
        if node.get("source_sha256"):
            result[filename] = node["source_sha256"]
        grouped[filename].append({
            "id": node["id"], "path": node.get("path"), "level": node["level"],
            "description": _original(node, "description"),
            "business_name": _original(node, "business_name"),
            "tags": _original(node, "tags"), "data_type": node.get("data_type"),
        })
    for filename, snapshots in grouped.items():
        result.setdefault(filename, sha256(json.dumps(sorted(snapshots, key=lambda x: x["id"]),
                          ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest())
    return result


def _review(reviews: dict, node_id: str) -> dict:
    data = reviews.get(node_id, {})
    return {"status": data.get("status", "Pending"), "completed": bool(data.get("completed", False)),
            "owner": _text(data.get("owner")), "comments": _text(data.get("comments")),
            "validated_at": _text(data.get("validated_at")), "pending_owner": _text(data.get("pending_owner")),
            "pending": _text(data.get("pending")), "final_comments": _text(data.get("final_comments")),
            "tag_comments": data.get("tag_comments", {})}


def _allowed(rule: dict) -> str:
    values = rule.get("values")
    if values:
        return ", ".join(_text(value) for value in values)
    if rule.get("external"):
        return MISSING_LIST
    return _text(rule.get("format")) or "Texto livre conforme orientação do catálogo"


def _category(rule: dict) -> str:
    return _text(rule.get("category")) or "Não identificada"


def _tag_proposal(tags: dict, tag: str) -> str:
    # Invalid native YAML booleans/numbers remain visible in the original column.
    # A blank proposal requires the reviewer to explicitly choose their text
    # replacement, avoiding an invisible automatic type conversion on reimport.
    value = tags.get(tag)
    return value if isinstance(value, str) else ""


def _issue_text(node: dict, nodes: list[dict], issues: list[dict]) -> str:
    ids = {node["id"]} | {child["id"] for child in nodes if child.get("object_id") == node["id"]}
    return "\n".join(dict.fromkeys(issue.get("message", "Pendência de revisão") for issue in issues if issue.get("id") in ids))


def _table_rows(nodes: list[dict], reviews: dict, issues: list[dict], references: dict | None) -> dict:
    catalog = rules.get_catalog(references)
    hashes = _source_hashes(nodes)
    sheets: dict[str, dict] = {
        "Objects": {"headers": ["Catálogo", "Schema / camada", "Nome técnico", "Tipo", "Nome de negócio atual", "Nome de negócio proposto", "Descrição atual", "Descrição proposta", "Domínio", "Subdomínio", "Finalidade", "Status da validação", "Validation completed", "Comentários do Data Owner", "Pendências técnicas", "ID estável", "Arquivo YAML de origem", "SHA256 da origem"], "rows": []},
        "Tags": {"headers": ["Caminho completo", "Nível", "Tag técnica", "Categoria", "Significado", "Presença no YAML", "Valor atual", "Valor proposto", "Valores permitidos / formato", "Situação", "Comentários", "ID estável", "Arquivo YAML de origem", "SHA256 da origem"], "rows": []},
        "Columns": {"headers": ["Catálogo", "Schema / camada", "Objeto", "Nome técnico da coluna", "Tipo técnico", "Descrição atual", "Descrição proposta"] + [label for tag in COLUMN_TAGS for label in (tag + " atual", tag + " proposto")] + ["Status da revisão", "Comentários", "ID estável", "Arquivo YAML de origem", "SHA256 da origem"], "rows": []},
        "Approval": {"headers": ["Data Owner", "Objeto", "Status final", "Validation completed", "Data da validação (declarada)", "Pendências técnicas", "Pendências do Owner", "Comentários finais", "Responsável pela pendência", "ID estável", "Arquivo YAML de origem", "SHA256 da origem"], "rows": []},
    }
    for node in nodes:
        node_id, level = node["id"], node["level"]
        original_tags = _original(node, "tags") or {}
        current_tags = node.get("tags") or {}
        rule_set = catalog.get(level, {})
        source = [_text(node_id), _text(node.get("filename")), hashes[node.get("filename", "")]]
        review = _review(reviews, node_id)
        if level in ("Table", "View"):
            values = [node.get("catalog"), node.get("schema"), node.get("name"), level,
                      _original(node, "business_name"), node.get("business_name"),
                      _original(node, "description"), node.get("description"),
                      current_tags.get("nm_dom"), current_tags.get("nm_subdom"), current_tags.get("nm_purpose"),
                      review["status"], "Sim" if review["completed"] else "Não", review["comments"],
                      _issue_text(node, nodes, issues)] + source
            sheets["Objects"]["rows"].append({"key": node_id, "node": node, "values": list(map(_text, values)),
                "editable": {5: ("business_name", None), 7: ("description", None), 11: ("review", "status"),
                             12: ("review", "completed"), 13: ("review", "comments")}})
            values = [review["owner"], node.get("path"), review["status"], "Sim" if review["completed"] else "Não",
                      review["validated_at"], _issue_text(node, nodes, issues), review["pending"], review["final_comments"],
                      review["pending_owner"]] + source
            sheets["Approval"]["rows"].append({"key": node_id, "node": node, "values": list(map(_text, values)),
                "editable": {0: ("review", "owner"), 2: ("review", "status"), 3: ("review", "completed"),
                             4: ("review", "validated_at"), 6: ("review", "pending"), 7: ("review", "final_comments"),
                             8: ("review", "pending_owner")}})
        if level == "Column":
            values = [node.get("catalog"), node.get("schema"), node.get("object_name") or ".".join(_text(node.get("path")).split(".")[-2:-1]),
                      node.get("name"), node.get("data_type"), _original(node, "description"), node.get("description")]
            editable = {6: ("description", None)}
            for tag in COLUMN_TAGS:
                editable[len(values) + 1] = ("tags", tag)
                values.extend([original_tags.get(tag), _tag_proposal(current_tags, tag)])
            editable[len(values)] = ("review", "status")
            editable[len(values) + 1] = ("review", "comments")
            values.extend([review["status"], review["comments"]] + source)
            sheets["Columns"]["rows"].append({"key": node_id, "node": node, "values": list(map(_text, values)), "editable": editable})
        for tag in sorted(set(rule_set) | set(original_tags) | set(current_tags)):
            rule = rule_set.get(tag, {})
            tag_issues = [issue.get("message", "Requer revisão") for issue in issues
                          if issue.get("id") == node_id and (issue.get("field") in (tag, "tags." + tag) or issue.get("tag") == tag)]
            if tag_issues:
                state = " | ".join(tag_issues)
            elif tag not in original_tags and tag not in current_tags:
                state = "Ausente (sem valor)"
            elif rule.get("external") and not rule.get("values"):
                state = MISSING_LIST
            elif not rule:
                state = "Tag sem regra oficial: requer revisão"
            elif current_tags.get(tag) != original_tags.get(tag):
                state = "Alterada: revalidar na aplicação"
            else:
                state = "Revisar conforme regras e pendências técnicas"
            values = [node.get("path"), level, tag, _category(rule), rule.get("description", "Tag preservada do YAML. Regra oficial não identificada."),
                      "Sim" if tag in original_tags else "Não", original_tags.get(tag), _tag_proposal(current_tags, tag),
                      _allowed(rule), state, review["tag_comments"].get(tag, "")] + source
            sheets["Tags"]["rows"].append({"key": (node_id, tag), "node": node, "tag": tag,
                                           "values": list(map(_text, values)), "editable": {7: ("tags", tag), 10: ("tag_comment", tag)}})
    return sheets


def _readme(workbook, styles: dict, references: dict | None):
    sheet = workbook.add_worksheet("Read Me")
    sheet.hide_gridlines(2)
    sheet.set_tab_color("#12324A")
    for col, width in enumerate([17, 30, 23, 65, 70, 30, 45]):
        sheet.set_column(col, col, width)
    sheet.merge_range("A1:G1", "Governança de dados | Guia do Data Owner", styles["title"])
    sheet.set_row(0, 36)
    blocks = [
        ("Objetivo", "Revise os metadados de negócio e registre decisões. Os arquivos YAML são processados em memória. Este Excel é um documento que você escolhe baixar e guardar; ele contém os metadados revisados."),
        ("Como preencher", "Edite somente as células em verde claro. As propostas já começam com o valor atual da sessão. Célula vazia mantém o valor existente. Para remover explicitamente um valor de negócio ou uma tag existente, escreva __DELETE__. A exclusão será revalidada na aplicação. Não exclua linhas, abas ou identificadores."),
        ("Abas e alterações", "Objects: nome e descrição de tabelas/views. Tags: valores em todos os níveis. Columns: descrição e sete tags usuais. Approval: responsável, conclusão e pendências. Quando a mesma informação aparece em duas abas, altere uma delas ou use o mesmo valor nas duas; alterações diferentes serão rejeitadas."),
        ("Hierarquia", "Catalog agrupa schemas. Schema organiza objetos. Table armazena dados. View representa uma consulta. Column é um atributo de tabela/view. A aprovação de um objeto não aprova automaticamente seus outros objetos, schema ou catálogo."),
        ("Bronze", "Camada normalmente associada aos dados recebidos da origem, preservando suas características originais."),
        ("Silver", "Camada normalmente associada aos dados tratados, padronizados e preparados para uso."),
        ("Gold", "Camada normalmente associada a dados consolidados e preparados para análises e consumo."),
        ("Data Product", "Schema para produtos de dados destinados a necessidades de negócio ou consumo específico."),
        ("Information Schema", "Área de metadados técnicos para consultar informações sobre os objetos do Databricks. As descrições das camadas são gerais; a implementação depende dos padrões internos da empresa."),
        ("Categorias", "[G] Governed Tag: valor controlado pela Governança. [T] Tag Livre: responsável define o texto respeitando o formato. [REDUNDANTE]: informação repetida de propósito em outro nível. Não identificada: o material disponível não atribui categoria individual à tag; ela não foi inferida."),
        ("Valores e listas", "Use os valores exatos das listas. Booleanos são textos true e false. Ausente não significa false nem nao_aplicavel. Lista oficial não carregada indica que a empresa precisa fornecer os valores oficiais de domínio, subdomínio, sistema ou área. Não foram criados valores para essas listas."),
        ("Tipos inválidos", "Se uma tag tiver tipo YAML inválido (como booleano nativo ou número), o valor original permanece visível e sua proposta começa vazia. Escolha explicitamente um texto permitido para corrigir o tipo. Deixar em branco mantém o original e sua pendência."),
        ("Status", "Pending: revisão ainda não iniciada. In review: revisão em andamento. Awaiting clarification: existe dúvida. Validated: revisão funcional aprovada. Validated with remarks: revisão concluída com pendências. Rejected: objeto reprovado. Use Comentários e Pendências do Owner para explicar as decisões."),
        ("Conclusão funcional", "Escolha Sim em Validation completed para cada objeto revisado. A conclusão do negócio não elimina inconsistências técnicas. Rejected continua reprovado, mesmo com revisão concluída. Somente objetos com status Validated podem integrar o pacote final; os demais são apresentados separadamente. Responsável e data são declarações do usuário, sem autenticação."),
        ("Exclusão explícita", "__DELETE__ é reservado para excluir valores existentes em campos de negócio, tags e comentários editáveis. Não se aplica a IDs, nomes técnicos, origens, status ou Validation completed. Exclusão que introduza inconsistência técnica será rejeitada. Células vazias, inclusive comentários, não apagam informações."),
        ("Reimportação", "Reimporte este Excel na aplicação com os mesmos YAMLs de origem. Confira mudanças, conflitos e rejeições antes de aplicar. IDs estáveis, nomes de origem e SHA256 detectam planilhas de outra versão. Fórmulas não são aceitas. O Excel não atualiza regras entre abas enquanto você edita; a validação final ocorre na aplicação."),
        ("Regras de consistência", "Mascaramento de uma coluna exige fl_mask=true no objeto. fl_apt_cons=true exige qualidade aprovado e data de aprovação. A classificação requer considerar as colunas e os níveis superiores; casos sem informação suficiente ficam pendentes. Valores e formatos precisam respeitar o catálogo abaixo."),
        ("Referências", "O catálogo abaixo usa a transcrição funcional fornecida. O PDF de referência apresenta ct_risco* em View como indicação de nota; a chave canônica usada é ct_risco. Categorias individuais não foram confirmadas. Novas evidências oficiais devem ser reconciliadas antes de alterar o catálogo."),
    ]
    row = 2
    for title, body in blocks:
        sheet.write_string(row, 0, title, styles["label"])
        sheet.merge_range(row, 1, row, 6, body, styles["body"])
        sheet.set_row(row, 48 if len(body) > 300 else 34)
        row += 1
    row += 1
    headers = ["Nível", "Tag", "Categoria", "Significado / propósito", "Valores permitidos", "Formato", "Observação"]
    for col, value in enumerate(headers):
        sheet.write_string(row, col, value, styles["header"])
    first = row
    for level, tags in rules.get_catalog(references).items():
        for tag, rule in tags.items():
            row += 1
            vals = [level, tag, _category(rule), rule.get("description"), _allowed(rule), rule.get("format"),
                    MISSING_LIST if rule.get("external") and not rule.get("values") else "Aplicar somente no nível indicado"]
            for col, value in enumerate(vals):
                sheet.write_string(row, col, _text(value), styles["locked"])
            sheet.set_row(row, 42)
    sheet.freeze_panes(first + 1, 2)
    sheet.autofilter(first, 0, row, 6)
    sheet.protect("", {"autofilter": True, "sort": True})


def export_workbook(nodes: list[dict], reviews: dict, issues: list[dict], references: dict | None = None) -> bytes:
    """Return the complete .xlsx as bytes without creating temporary files."""
    buffer = BytesIO()
    workbook = xlsxwriter.Workbook(buffer, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False})
    workbook.set_properties({"title": "Validação de metadados por Data Owners", "author": "Governança de dados", "comments": "Processamento em memória"})
    common = {"font_name": "Aptos", "font_size": 11, "valign": "top", "text_wrap": True}
    styles = {
        "title": workbook.add_format({**common, "font_size": 20, "bold": True, "font_color": "#FFFFFF", "bg_color": "#12324A", "valign": "vcenter"}),
        "header": workbook.add_format({**common, "font_color": "#FFFFFF", "bg_color": "#156C72", "bold": True, "valign": "vcenter"}),
        "locked": workbook.add_format({**common, "bg_color": "#F1F5F8", "font_color": "#243B4B", "locked": True, "bottom": 1, "bottom_color": "#E2E9EE"}),
        "editable": workbook.add_format({**common, "bg_color": "#E3F5EF", "font_color": "#123C32", "locked": False, "bottom": 1, "bottom_color": "#C4E6D9"}),
        "body": workbook.add_format({**common, "font_color": "#243B4B"}),
        "label": workbook.add_format({**common, "bold": True, "font_color": "#156C72"}),
        "warning": workbook.add_format({"bg_color": "#FFF1D6", "font_color": "#8B5415"}),
        "done": workbook.add_format({"bg_color": "#CCEDE0", "font_color": "#18533D"}),
    }
    _readme(workbook, styles, references)
    catalog = rules.get_catalog(references)
    tables = _table_rows(nodes, reviews, issues, references)
    worksheet_map = {name: workbook.add_worksheet(name) for name in tables}
    list_sheet = workbook.add_worksheet("_Listas")
    named_lists: dict[tuple, str] = {}

    def named(values: list) -> str:
        key = tuple(_text(v) for v in values)
        if key not in named_lists:
            col = len(named_lists)
            name = f"GovList_{col + 1:04d}"
            for row, value in enumerate(key):
                list_sheet.write_string(row, col, value)
            ref_col = xl_col_to_name(col)
            workbook.define_name(name, f"='_Listas'!${ref_col}$1:${ref_col}${len(key)}")
            named_lists[key] = name
        return "=" + named_lists[key]

    for sheet_name, table in tables.items():
        sheet = worksheet_map[sheet_name]
        headers, rows = table["headers"], table["rows"]
        sheet.hide_gridlines(2)
        sheet.set_tab_color("#156C72")
        sheet.set_zoom(85)
        sheet.merge_range(0, 0, 0, len(headers) - 1, "Governança de dados | " + sheet_name, styles["title"])
        sheet.set_row(0, 36)
        sheet.merge_range(1, 0, 1, len(headers) - 1, "Preencha as células verdes. Não remova linhas nem altere IDs/origens. Reimporte para recalcular regras e conferir mudanças.", styles["body"])
        sheet.set_row(1, 28)
        sheet.set_row(HEADER_ROW, 44)
        for col, header in enumerate(headers):
            width = 42 if any(word in header for word in ("Descrição", "Comentários", "Pendências", "Significado", "Valores permitidos")) else 24
            if "SHA256" in header:
                width = 24
            sheet.set_column(col, col, width)
            sheet.write_string(HEADER_ROW, col, header, styles["header"])
        for offset, record in enumerate(rows, HEADER_ROW + 1):
            sheet.set_row(offset, 46)
            for col, value in enumerate(record["values"]):
                editable = record["editable"].get(col)
                if len(value) > 32767:
                    raise ValueError("Um campo excede o limite de texto do Excel (32.767 caracteres).")
                sheet.write_string(offset, col, value, styles["editable" if editable else "locked"])
                if not editable:
                    continue
                field, tag = editable
                values = None
                if (field, tag) == ("review", "status"):
                    values = STATUSES
                elif (field, tag) == ("review", "completed"):
                    values = ["Não", "Sim"]
                elif field == "tags":
                    values = catalog.get(record["node"]["level"], {}).get(tag, {}).get("values")
                    if values and tag in record["node"].get("tags", {}):
                        values = list(values) + [DELETE_TOKEN]
                if values:
                    sheet.data_validation(offset, col, offset, col, {"validate": "list", "source": named(values),
                        "ignore_blank": True, "error_type": "stop", "error_title": "Valor fora da lista", "error_message": "Selecione um valor oficial na lista.",
                        "input_title": "Campo de revisão", "input_message": "Em branco mantém. __DELETE__ remove um valor existente quando permitido."})
                if field == "review" and tag == "status":
                    sheet.conditional_format(offset, col, offset, col, {"type": "text", "criteria": "containing", "value": "remarks", "format": styles["warning"]})
                    sheet.conditional_format(offset, col, offset, col, {"type": "cell", "criteria": "==", "value": '"Validated"', "format": styles["done"]})
                    sheet.conditional_format(offset, col, offset, col, {"type": "cell", "criteria": "==", "value": '"Rejected"', "format": styles["warning"]})
        sheet.autofilter(HEADER_ROW, 0, HEADER_ROW + max(1, len(rows)), len(headers) - 1)
        sheet.freeze_panes(HEADER_ROW + 1, min(4, len(headers) - 1))
        sheet.protect("", {"autofilter": True, "sort": True, "select_locked_cells": True, "select_unlocked_cells": True})
        sheet.set_landscape()
        sheet.fit_to_pages(1, 0)
        sheet.repeat_rows(HEADER_ROW)
    # The visible order remains exactly the five requested tabs.
    list_sheet.hide()
    workbook.close()
    return buffer.getvalue()


def _inspect_container(content: bytes):
    if len(content) > MAX_XLSX_BYTES:
        raise ValueError("O Excel excede o limite de 20 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 1000 or sum(info.file_size for info in infos) > MAX_XML_BYTES:
                raise ValueError("Excel muito grande após descompactação.")
            if len({info.filename for info in infos}) != len(infos):
                raise ValueError("Excel possui entradas ZIP duplicadas.")
            for info in infos:
                path = PurePosixPath(info.filename)
                if path.is_absolute() or ".." in path.parts or "\\" in info.filename:
                    raise ValueError("Estrutura ZIP do Excel inválida.")
                if info.flag_bits & 1 or info.file_size > 30 * 1024 * 1024 or (info.file_size > 1024 * 1024 and info.file_size / max(1, info.compress_size) > 400):
                    raise ValueError("Excel criptografado ou com expansão excessiva.")
                lower = info.filename.lower()
                if any(value in lower for value in ("vbaproject", "externallinks/", "embeddings/")):
                    raise ValueError("Excel com macros, vínculos externos ou conteúdo incorporado não é aceito.")
                if lower.endswith((".xml", ".rels")):
                    payload = archive.read(info)
                    if re.search(br"<!\s*(DOCTYPE|ENTITY)\b", payload, re.IGNORECASE):
                        raise ValueError("Declarações XML externas não são aceitas.")
    except BadZipFile as exc:
        raise ValueError("Arquivo Excel inválido ou corrompido.") from exc


def _validate_proposal(node: dict, field: str, tag: str | None, value: Any, catalog: dict) -> str | None:
    if value is not None and not isinstance(value, str):
        return "O valor proposto deve ser texto."
    if len(value or "") > 32000:
        return "Texto proposto excede 32.000 caracteres."
    if field != "tags" or value is None:
        return None
    rule = catalog.get(node["level"], {}).get(tag)
    if not rule:
        return "Tag sem regra oficial; mantenha o valor e solicite definição à Governança."
    values = rule.get("values")
    if values and value not in values:
        return "Valor fora da lista oficial permitida."
    if rule.get("external") and not values and not (tag == "nm_sist_orig" and value == "multiplas_origens"):
        return MISSING_LIST + "; carregue a referência antes de alterar o valor."
    official_issues = rules.validate_tag(node["level"], tag, value)
    if not rule.get("external") and official_issues:
        return official_issues[0]["message"]
    fmt = _text(rule.get("format")).lower()
    if (tag == "dt_aprv_gov" or "yyyy-mm-dd" in fmt) and value:
        try:
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
                raise ValueError
            date.fromisoformat(value)
        except ValueError:
            return "Data inválida. Use YYYY-MM-DD."
    if (tag == "ct_prod_neg" or "snake_case" in fmt) and value and not re.fullmatch(r"[a-z][a-z0-9]*(?:_[a-z0-9]+)*", value):
        return "Use snake_case com letras minúsculas, números e sublinhados."
    if (tag in ("nm_own_stwd",) or (tag == "nm_cust" and node["level"] in ("Table", "View")) or fmt in ("email", "e-mail")) and value and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value):
        return "Informe um e-mail válido."
    if tag in ("grp_owner", "grp_own_stwd") and value:
        suffix = "owners" if tag == "grp_owner" else "stewards"
        if not re.fullmatch(r"grp-dbx-data-[a-z0-9]+(?:-[a-z0-9]+)*-" + suffix, value):
            return "Grupo fora do formato grp-dbx-data-{dominio}-" + suffix + "."
    return None


def preview_workbook(content: bytes, nodes: list[dict], reviews: dict, references: dict | None = None) -> dict:
    """Check a review workbook and preview safe edits without mutating the session."""
    result = {"edits": [], "reviews": deepcopy(reviews), "changes": [], "rejected": [], "missing": []}

    def reject(message: str, **details):
        result["rejected"].append({"message": message, **details})

    try:
        _inspect_container(content)
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        reject(str(exc) if isinstance(exc, ValueError) else "Não foi possível ler o Excel. Confirme o formato .xlsx.", code="invalid_workbook")
        return result
    try:
        if set(workbook.sheetnames) - set(VISIBLE_SHEETS + ["_Listas"]):
            reject("O Excel contém abas desconhecidas.", code="unknown_sheets")
            return result
        for name in VISIBLE_SHEETS:
            if name not in workbook.sheetnames:
                result["missing"].append({"sheet": name, "message": "Aba obrigatória ausente."})
        if result["missing"]:
            return result
        for sheet in workbook:
            if sheet.max_row > MAX_ROWS or sheet.max_column > 128:
                reject("Aba excede o limite de linhas ou colunas.", sheet=sheet.title, code="sheet_limits")
                return result
            for row in sheet.iter_rows():
                if any(cell.data_type == "f" for cell in row):
                    reject("Fórmulas não são aceitas no Excel de revisão.", sheet=sheet.title, row=row[0].row, code="formula")
        if result["rejected"]:
            return result
        tables = _table_rows(nodes, reviews, rules.validate_nodes(nodes, references), references)
        node_map = {node["id"]: node for node in nodes}
        candidates: dict[tuple, list] = defaultdict(list)
        review_candidates: dict[tuple, list] = defaultdict(list)
        comment_candidates: dict[tuple, list] = defaultdict(list)
        blocked: set[tuple] = set()

        def block_record(record):
            for field, tag in record["editable"].values():
                blocked.add((record["node"]["id"], field, tag))
        for name, table in tables.items():
            sheet = workbook[name]
            headers = table["headers"]
            actual_headers = [_text(cell.value) for cell in next(sheet.iter_rows(min_row=HEADER_ROW + 1, max_row=HEADER_ROW + 1, max_col=len(headers)))]
            if actual_headers != headers:
                reject("Cabeçalhos alterados. Use o modelo exportado pela aplicação.", sheet=name, code="headers")
                for record in table["rows"]:
                    block_record(record)
                continue
            expected = {record["key"]: record for record in table["rows"]}
            id_col = headers.index("ID estável")
            tag_col = headers.index("Tag técnica") if name == "Tags" else None
            incoming = []
            for cells in sheet.iter_rows(min_row=HEADER_ROW + 2, max_col=max(len(headers), sheet.max_column)):
                if not any(cell.value is not None for cell in cells):
                    continue
                values = [_text(cell.value) for cell in cells]
                node_id = values[id_col]
                key = (node_id, values[tag_col]) if tag_col is not None else node_id
                incoming.append((key, cells[0].row, values, cells))
            duplicates = {key for key, count in Counter(key for key, *_ in incoming).items() if count > 1}
            found = {key for key, *_ in incoming}
            for key in expected.keys() - found:
                record = expected[key]
                result["missing"].append({"sheet": name, "id": record["node"]["id"], "tag": record.get("tag"), "message": "Registro esperado ausente no Excel."})
                block_record(record)
            for key, row_num, values, cells in incoming:
                node_id = key[0] if isinstance(key, tuple) else key
                details = {"sheet": name, "row": row_num, "id": node_id}
                if key in duplicates:
                    reject("Chave duplicada; nenhuma alteração deste registro foi aplicada.", **details, code="duplicate")
                    if key in expected:
                        block_record(expected[key])
                    continue
                record = expected.get(key)
                if not record:
                    reject("Registro desconhecido ou identificador alterado.", **details, code="unknown_row")
                    continue
                if any(values[len(headers):]):
                    reject("Há conteúdo em colunas não previstas no modelo.", **details, code="extra_columns")
                    block_record(record)
                    continue
                if any(values[col] != old for col, old in enumerate(record["values"]) if col not in record["editable"]):
                    reject("Identidade, origem ou valor original alterado. A planilha pode pertencer a outra versão dos YAMLs.", **details, code="immutable")
                    block_record(record)
                    continue
                for col, (field, tag) in record["editable"].items():
                    value = values[col]
                    # A blank cell means keep. Deletion is a separate, explicit
                    # instruction; a cleared cell must never erase source data.
                    if not value.strip():
                        continue
                    if value == record["values"][col]:
                        continue
                    details_field = {**details, "field": field, "tag": tag}
                    # Actual numeric/boolean cells must not silently coerce into tag strings.
                    raw_value = cells[col].value
                    if raw_value is not None and not isinstance(raw_value, str):
                        reject("Preencha o campo como texto, usando os valores exatos do catálogo.", **details_field, code="value_type")
                        blocked.add((node_id, field, tag))
                        continue
                    if field == "review":
                        if value == DELETE_TOKEN and tag not in ("status", "completed"):
                            if not record["values"][col]:
                                reject("O campo já está ausente; não há valor para excluir.", **details_field, code="deletion_absent")
                                blocked.add((node_id, field, tag))
                                continue
                            value = ""
                        if tag == "status" and value not in STATUSES:
                            reject("Status de revisão inválido.", **details_field, code="status")
                            blocked.update({(node_id, "review", "status"), (node_id, "review", "completed")})
                            continue
                        if tag == "completed" and value not in ("Sim", "Não"):
                            reject("Conclusão deve ser Sim ou Não.", **details_field, code="completed")
                            blocked.update({(node_id, "review", "status"), (node_id, "review", "completed")})
                            continue
                        if tag == "validated_at" and value:
                            try:
                                datetime.fromisoformat(value.replace("Z", "+00:00"))
                            except ValueError:
                                reject("Data/hora declarada inválida. Use formato ISO 8601.", **details_field, code="datetime")
                                blocked.add((node_id, field, tag))
                                continue
                        review_candidates[(node_id, tag)].append((value == "Sim" if tag == "completed" else value, details_field))
                    elif field == "tag_comment":
                        if value == DELETE_TOKEN:
                            if not record["values"][col]:
                                reject("O comentário já está ausente.", **details_field, code="deletion_absent")
                                blocked.add((node_id, field, tag))
                                continue
                            value = ""
                        comment_candidates[(node_id, tag)].append((value, details_field))
                    else:
                        if value == DELETE_TOKEN:
                            existing = tag in record["node"].get("tags", {}) if field == "tags" else record["node"].get(field) is not None
                            if not existing:
                                reject("O campo já está ausente; não há valor para excluir.", **details_field, code="deletion_absent")
                                blocked.add((node_id, field, tag))
                                continue
                            value = None
                        error = _validate_proposal(record["node"], field, tag, value, rules.get_catalog(references))
                        if error:
                            reject(error, **details_field, code="invalid_proposal")
                            blocked.add((node_id, field, tag))
                            continue
                        candidates[(node_id, field, tag)].append((value, details_field))
        for (node_id, field, tag), proposals in candidates.items():
            if (node_id, field, tag) in blocked:
                reject("Alteração bloqueada porque o registro correspondente está ausente, duplicado ou foi alterado.", id=node_id, field=field, tag=tag, code="blocked_record")
                continue
            unique = {value for value, _ in proposals}
            if len(unique) != 1:
                reject("Propostas conflitantes em abas diferentes; altere uma aba ou use o mesmo valor nas duas.", id=node_id, field=field, tag=tag, code="conflict")
                continue
            value = proposals[0][0]
            edit = {"id": node_id, "field": field, "value": value}
            if tag:
                edit["tag"] = tag
            result["edits"].append(edit)
            before = node_map[node_id].get("tags", {}).get(tag) if field == "tags" else node_map[node_id].get(field)
            result["changes"].append({**edit, "before": before, "after": value})
        accepted_review_keys = set()
        for (node_id, field), proposals in review_candidates.items():
            if (node_id, "review", field) in blocked:
                reject("Revisão bloqueada porque o registro correspondente está ausente, duplicado ou foi alterado.", id=node_id, field=field, code="blocked_record")
                continue
            if len({value for value, _ in proposals}) != 1:
                reject("Informações de aprovação conflitantes entre abas.", id=node_id, field=field, code="review_conflict")
                continue
            value = proposals[0][0]
            accepted_review_keys.add((node_id, field))
            result["reviews"].setdefault(node_id, {}).update({field: value})
            result["changes"].append({"id": node_id, "field": "review." + field, "before": _review(reviews, node_id).get(field), "after": value})
        for (node_id, tag), proposals in comment_candidates.items():
            if (node_id, "tag_comment", tag) in blocked:
                continue
            value = proposals[0][0]
            result["reviews"].setdefault(node_id, {}).setdefault("tag_comments", {})[tag] = value
            result["changes"].append({"id": node_id, "field": "tag_comments", "tag": tag, "after": value})
        # Evaluate all proposals together, allowing jointly consistent corrections.
        # New cross-object errors reject tag proposals for the affected object and
        # its columns; descriptions and independent objects can still be reviewed.
        proposed = deepcopy(nodes)
        proposed_map = {node["id"]: node for node in proposed}
        for edit in result["edits"]:
            node = proposed_map[edit["id"]]
            if edit["field"] == "tags":
                if edit["value"] is None:
                    node.setdefault("tags", {}).pop(edit["tag"], None)
                else:
                    node.setdefault("tags", {})[edit["tag"]] = edit["value"]
            else:
                node[edit["field"]] = edit["value"]
        result["issues"] = rules.validate_nodes(proposed, references)
        existing_errors = {(i.get("id"), i.get("field"), i.get("code")) for i in rules.validate_nodes(nodes, references) if i.get("severity") == "error"}
        new_errors = [i for i in result["issues"] if i.get("severity") == "error" and (i.get("id"), i.get("field"), i.get("code")) not in existing_errors]
        rejected_groups = set()
        for issue in new_errors:
            issue_node = node_map.get(issue.get("id"), {})
            rejected_groups.add(issue_node.get("object_id") or issue.get("id"))
        if rejected_groups:
            accepted = []
            for edit in result["edits"]:
                node = node_map[edit["id"]]
                group = node.get("object_id") or node["id"]
                if edit["field"] == "tags" and group in rejected_groups:
                    reject("A proposta introduz inconsistência técnica no objeto ou em suas colunas. Revise as propostas em conjunto.",
                           id=edit["id"], field=edit["field"], tag=edit.get("tag"), code="new_consistency_error")
                    original_tags = node.get("tags", {})
                    target = proposed_map[node["id"]].setdefault("tags", {})
                    if edit["tag"] in original_tags:
                        target[edit["tag"]] = original_tags[edit["tag"]]
                    else:
                        target.pop(edit["tag"], None)
                else:
                    accepted.append(edit)
            result["edits"] = accepted
            accepted_keys = {(e["id"], e["field"], e.get("tag")) for e in accepted}
            result["changes"] = [change for change in result["changes"] if change.get("field") not in ("tags", "description", "business_name")
                                 or (change["id"], change["field"], change.get("tag")) in accepted_keys]
            result["issues"] = rules.validate_nodes(proposed, references)
        for node_id, review in result["reviews"].items():
            if node_id not in proposed_map:
                continue
            completion_changed = (node_id, "completed") in accepted_review_keys
            status_changed = (node_id, "status") in accepted_review_keys
            if status_changed and not completion_changed and review.get("status") in ("Validated", "Validated with remarks"):
                review["completed"] = True
            if completion_changed and not review.get("completed") and review.get("status") in ("Validated", "Validated with remarks"):
                review["status"] = "In review"
            if review.get("completed") and review.get("status") not in ("Rejected", "Awaiting clarification"):
                pending = _issue_text(proposed_map[node_id], proposed, result["issues"])
                if pending or _text(review.get("pending")).strip():
                    review["status"] = "Validated with remarks"
                elif review.get("status", "Pending") not in ("Validated", "Validated with remarks"):
                    review["status"] = "Validated"
            elif review.get("status") in ("Validated", "Validated with remarks"):
                review["status"] = "In review"
        return result
    finally:
        workbook.close()
