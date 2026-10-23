"""Planilhas portáteis de revisão: todo o processamento permanece em memória.

As fontes ficam dentro do próprio Excel entregue ao usuário. A aba veryHidden
evita alterações acidentais; ela não é criptografia ou assinatura digital.
"""
from __future__ import annotations

import base64
import binascii
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
import json
import math
import posixpath
import re
from zipfile import BadZipFile, ZipFile
import zlib
from defusedxml import ElementTree

import openpyxl
import xlsxwriter
from xlsxwriter.utility import xl_col_to_name

from . import engine, rules


VISIBLE_SHEETS = ["Comece aqui", "Objetos", "Colunas", "Tags"]
PAYLOAD_SHEET = "_metadata"
LIST_SHEET = "_listas"
HEADER_ROW = 3  # Índice zero: cabeçalho na linha 4 do Excel.
DELETE_TOKEN = "[REMOVER]"
LEGACY_VERSION = "metadata-reviewer-owner-v1"
PREVIOUS_VERSION = "metadata-reviewer-owner-v2"
VERSION = "metadata-reviewer-owner-v3"
SUPPORTED_VERSIONS = frozenset((LEGACY_VERSION, PREVIOUS_VERSION, VERSION))
MAX_XLSX_BYTES = 30 * 1024 * 1024
MAX_XML_BYTES = 160 * 1024 * 1024
MAX_PAYLOAD_BYTES = engine.MAX_BATCH_BYTES * 4 // 3 + 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 256
MAX_CELLS = 2_000_000
CHUNK_SIZE = 30_000
COLUMN_TAGS = ["fl_reg_apl", "fl_atr_pii", "nm_atr_pii", "fl_atr_sens",
               "fl_atr_mask", "nm_atr_class_inf", "fl_metadata_ai_generate"]
LEVEL_NAMES = {"Catalog": "Catálogo", "Schema": "Schema", "Table": "Tabela",
               "View": "Visão", "Column": "Coluna"}
TAG_LABELS = {
    "nm_environment": "Ambiente", "dt_aprv_gov": "Aprovação da governança",
    "nm_cust": "Custodiante", "nm_dom": "Domínio", "fl_reg_apl": "Requisito regulatório",
    "ct_risco": "Risco", "nm_class_inf": "Classificação da informação",
    "ct_area_neg": "Área de negócio", "grp_owner": "Grupo de Owners",
    "fl_exempt": "Isenção de governança", "nm_purpose": "Finalidade",
    "grp_own_stwd": "Grupo de Stewards", "fl_ropa": "Registro de tratamento",
    "fl_metadata_ai_generate": "Metadado gerado por IA", "nm_cam": "Camada de dados",
    "nm_lgl_lgpd": "Base legal LGPD", "nm_own_stwd": "Responsável de negócio",
    "nm_subdom": "Subdomínio", "nm_sist_orig": "Sistema de origem",
    "fl_pii": "Dados pessoais", "fl_sens": "Dados sensíveis", "nm_retencao": "Retenção",
    "fl_mask": "Mascaramento", "ct_qld_status": "Qualidade dos dados",
    "fl_apt_cons": "Apto para consumo", "ct_cons_exec": "Consumo executivo",
    "fl_legado": "Objeto legado", "ct_prod_neg": "Produto de negócio",
    "fl_atr_pii": "Dado pessoal", "nm_atr_pii": "Tipo de dado pessoal",
    "fl_atr_sens": "Dado sensível", "fl_atr_mask": "Dado mascarado",
    "nm_atr_class_inf": "Classificação da informação",
}


def _text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _cell_text(value):
    """Excel has a hard character limit; the source payload remains complete."""
    text = _text(value)
    if len(text) > 32_000:
        return text[:31_900] + "\n[Texto extenso: conteúdo completo preservado no YAML de origem.]"
    return text


def _tag_text(value, rule):
    if isinstance(value, str) and rule.get("values") == rules.BOOL:
        return {"true": "Sim", "false": "Não"}.get(value, _cell_text(value))
    return _cell_text(value)


def _source(files):
    if not isinstance(files, list) or not files:
        raise ValueError("A planilha precisa conter pelo menos um YAML de origem.")
    batch = engine.parse_batch(files)
    if batch.errors or batch.duplicates or not batch.nodes:
        raise ValueError("Os YAMLs de origem estão inválidos, incompletos ou contêm objetos duplicados.")
    return batch


def _allowed(rule):
    if rule.get("values"):
        if rule["values"] == rules.BOOL:
            return "Sim, Não"
        return ", ".join(rule["values"])
    if rule.get("external"):
        return "Texto; lista oficial indisponível. Confirmar com a governança."
    if rule.get("format"):
        return "Formato: " + rule["format"]
    return "Texto livre; tag sem regra documentada. Confirmar com a governança."


def _tables(nodes, include_column_tags=True):
    """The same deterministic schema drives authoring and independent validation."""
    catalog = rules.get_catalog()
    extra_column_tags = sorted({tag for node in nodes if node["level"] == "Column"
                                for tag in node["tags"] if tag not in COLUMN_TAGS})
    column_tags = COLUMN_TAGS + extra_column_tags
    hidden = ["ID interno", "Arquivo de origem", "SHA256 de origem"]
    tables = {
        "Objetos": {"headers": ["Tipo de objeto", "Catálogo", "Schema", "Nome técnico",
            "Nome de negócio atual", "Nome de negócio proposto", "Descrição atual",
            "Descrição proposta", "Comentário para a governança"] + hidden,
            "rows": [], "widths": [16, 24, 24, 30, 34, 34, 65, 65, 48], "hidden_start": 9},
        "Colunas": {"headers": ["Catálogo", "Schema", "Tabela / visão", "Coluna",
            "Tipo de dado", "Descrição atual", "Descrição proposta"] +
            [label for tag in column_tags for label in
             (TAG_LABELS.get(tag, "Tag: " + tag) + " — atual",
              TAG_LABELS.get(tag, "Tag: " + tag) + " — proposto")] +
            ["Comentário para a governança"] + hidden,
            "rows": [], "widths": [22, 22, 40, 30, 18, 64, 64] +
            [21 if catalog["Column"].get(tag, {}).get("values") == rules.BOOL else 27 for tag in column_tags for _ in range(2)] + [48],
            "hidden_start": 8 + 2 * len(column_tags), "column_tags": column_tags},
        "Tags": {"headers": ["Tipo de objeto", "Objeto", "Informação de governança", "O que significa",
            "Valor atual", "Valor proposto", "Valores aceitos / orientação",
            "Comentário para a governança", "Tag técnica"] + hidden,
            "rows": [], "widths": [16, 56, 30, 60, 38, 38, 65, 48, 28], "hidden_start": 9},
    }
    if len(tables["Colunas"]["headers"]) > MAX_COLUMNS:
        raise ValueError("Há tags de coluna em excesso; divida o lote em planilhas menores.")
    for node in nodes:
        tags, level, nid = node["tags"], node["level"], node["id"]
        origin = [nid, node["filename"], node["source_sha256"]]
        if level != "Column":
            values = [LEVEL_NAMES[level], node["catalog"], node["schema"], node["name"],
                      node.get("business_name"), "", node.get("description"), "", ""] + origin
            tables["Objetos"]["rows"].append({"key": nid, "node": node,
                "values": list(map(_cell_text, values)), "editable": {5: ("business_name", None),
                7: ("description", None), 8: ("note", None)}})
            # Only existing tags: an Owner can edit metadata without creating fields.
            for tag in tags:
                rule = catalog.get(level, {}).get(tag, {})
                values = [LEVEL_NAMES[level], node["path"], TAG_LABELS.get(tag, "Tag: " + tag),
                          rule.get("description", "Tag existente sem significado documentado. Confirmar com a governança."),
                          _tag_text(tags[tag], rule), "", _allowed(rule), "", tag] + origin
                tables["Tags"]["rows"].append({"key": (nid, tag), "node": node, "tag": tag,
                    "values": list(map(_cell_text, values)), "editable": {5: ("tags", tag), 7: ("note", tag)}})
        else:
            parent_name = node["path"].rsplit(".", 1)[0]
            values = [node["catalog"], node["schema"], parent_name, node["name"],
                      node.get("data_type", node.get("type")), node.get("description"), ""]
            editable = {6: ("description", None)}
            for tag in column_tags:
                col = len(values)
                values.extend([_tag_text(tags.get(tag), catalog["Column"].get(tag, {})), ""])
                if tag in tags:
                    editable[col + 1] = ("tags", tag)
            editable[len(values)] = ("note", None)
            values.extend([""] + origin)
            tables["Colunas"]["rows"].append({"key": nid, "node": node,
                "values": list(map(_cell_text, values)), "editable": editable})
            if include_column_tags:
                for tag, value in tags.items():
                    rule = catalog["Column"].get(tag, {})
                    values = [LEVEL_NAMES[level], node["path"], TAG_LABELS.get(tag, "Tag: " + tag),
                              rule.get("description", "Tag existente sem significado documentado. Confirmar com a governança."),
                              _tag_text(value, rule), "", _allowed(rule), "", tag] + origin
                    tables["Tags"]["rows"].append({"key": (nid, tag), "node": node, "tag": tag,
                        "values": list(map(_cell_text, values)), "editable": {5: ("tags", tag), 7: ("note", tag)}})
    for table in tables.values():
        if len(table["rows"]) > MAX_ROWS - HEADER_ROW - 1:
            raise ValueError("Há linhas em excesso; divida o lote em planilhas menores.")
    if sum(len(t["rows"]) * len(t["headers"]) for t in tables.values()) > MAX_CELLS:
        raise ValueError("O lote excede o limite de células; divida em planilhas menores.")
    return tables


def _styles(book):
    common = {"font_name": "Calibri", "font_size": 11, "valign": "top", "text_wrap": True,
              "border": 1, "border_color": "#DEE5ED", "num_format": "@"}
    return {
        "title": book.add_format({"font_name": "Calibri", "font_size": 19, "bold": True,
                                  "font_color": "#16324F", "valign": "vcenter"}),
        "hint": book.add_format({"font_name": "Calibri", "font_size": 11, "font_color": "#41546A",
                                 "text_wrap": True, "valign": "top"}),
        "header": book.add_format({**common, "bold": True, "bg_color": "#16324F", "font_color": "white"}),
        "current": book.add_format({**common, "bg_color": "#F3F6F9", "font_color": "#34465A"}),
        "input": book.add_format({**common, "bg_color": "#FFF2C6", "locked": False}),
        "note": book.add_format({**common, "bg_color": "#EAF3FC", "locked": False}),
        "inactive": book.add_format({**common, "bg_color": "#E5E9EE", "font_color": "#6A7582"}),
    }


def _guide(book, styles, nodes, issues):
    ws = book.add_worksheet("Comece aqui")
    ws.hide_gridlines(2)
    ws.set_tab_color("#16324F")
    ws.set_column(0, 0, 4)
    ws.set_column(1, 1, 27)
    ws.set_column(2, 5, 24)
    lines = [
        ("Revisão de metadados", "title", 34),
        ("1. Leia Objetos e Colunas. Confira os nomes e as descrições. Use Tags para as informações de governança de todos os níveis, inclusive as colunas.", "hint", 45),
        ("Na aba Tags, filtre Tipo de objeto = Coluna para revisar as tags colunares. O campo Objeto mostra o caminho completo, com a tabela e a coluna.", "hint", 45),
        ("Uma tag de coluna pode ser proposta em Tags ou Colunas. Se preencher nas duas abas, informe o mesmo valor. Propostas diferentes precisam ser corrigidas antes da conversão.", "hint", 50),
        ("2. Escreva somente nos campos amarelos (propostas) e azuis (comentários). Use os filtros para localizar a tabela ou coluna desejada.", "hint", 45),
        ("3. Salve este Excel em .xlsx e envie à governança. Ela carregará apenas este arquivo na aba Excel → YAML do aplicativo.", "hint", 45),
        ("Célula proposta em branco = manter o valor original. Para excluir um valor existente, escreva exatamente [REMOVER]. Os nomes técnicos e tipos de dados ficam protegidos.", "hint", 55),
        ("Tags ausentes no YAML ficam cinzas e não são editáveis. Não insira nem exclua linhas, colunas ou abas. Para ordenar, selecione a linha inteira, incluindo as colunas ocultas.", "hint", 55),
        ("Nas listas, Sim e Não geram os textos true e false no YAML. Datas usam AAAA-MM-DD. Outros valores seguem as opções apresentadas na planilha.", "hint", 50),
        ("Preencha com texto ou valores, sem fórmulas. Se usar uma fórmula para preparar a resposta, copie o resultado e use Colar especial > Valores antes de enviar o Excel.", "hint", 48),
        ("Comentários são orientações para a governança e não viram campos do YAML. Não há status de aprovação obrigatório neste arquivo.", "hint", 42),
        ("Este Excel inclui os YAMLs de origem em uma aba técnica oculta, para reconstruir os arquivos sem pedir novo upload. Aba oculta não é criptografia. Cuide deste Excel como dos metadados corporativos.", "hint", 60),
        ("Descrições extensas podem aparecer resumidas na célula. O conteúdo completo está preservado na origem; uma proposta em branco sempre o mantém inteiro.", "hint", 42),
        (f"Conteúdo: {sum(n['level'] != 'Column' for n in nodes)} objeto(s), {sum(n['level'] == 'Column' for n in nodes)} coluna(s). O aplicativo processa os arquivos em memória e não persiste os uploads.", "hint", 44),
    ]
    row = 1
    for text, style, height in lines:
        ws.merge_range(row, 1, row, 5, text, styles[style])
        ws.set_row(row, height)
        row += 1
    advisory = list(dict.fromkeys(i.get("message", "") for i in issues if i.get("message")))
    if advisory:
        ws.merge_range(row + 1, 1, row + 1, 5, "Observações da origem (não impedem a conversão)", styles["header"])
        for message in advisory[:30]:
            row += 1
            ws.merge_range(row + 1, 1, row + 1, 5, message, styles["hint"])
            ws.set_row(row + 1, 36)
    row += 4
    ws.merge_range(row, 1, row, 5, "Como ler as informações de Colunas", styles["header"])
    for tag in COLUMN_TAGS:
        row += 1
        rule = rules.get_catalog()["Column"][tag]
        ws.merge_range(row, 1, row, 5, f"{TAG_LABELS[tag]} ({tag}): {rule['description']} Valores: {_allowed(rule)}", styles["hint"])
        ws.set_row(row, 47)
    ws.freeze_panes(2, 0)
    ws.protect("", {"select_locked_cells": True})


def _write_payload(book, files):
    payload = json.dumps({"version": VERSION, "files": files}, ensure_ascii=False,
                         sort_keys=True, separators=(",", ":")).encode("utf-8")
    if len(payload) > MAX_PAYLOAD_BYTES:
        raise ValueError("O lote excede o tamanho máximo para o Excel portátil; divida os arquivos.")
    packed = base64.b64encode(zlib.compress(payload, 6)).decode("ascii")
    chunks = [packed[i:i + CHUNK_SIZE] for i in range(0, len(packed), CHUNK_SIZE)]
    manifest = {"version": VERSION, "sha256": sha256(payload).hexdigest(),
                "length": len(payload), "chunks": len(chunks), "encoding": "zlib+base64"}
    ws = book.add_worksheet(PAYLOAD_SHEET)
    ws.write_string(0, 0, json.dumps(manifest, sort_keys=True))
    for index, chunk in enumerate(chunks, 1):
        ws.write_string(index, 0, chunk)
    ws.very_hidden()
    ws.protect()


def build_owner_workbook(files: list[dict], nodes: list[dict], issues: list[dict]) -> bytes:
    """Create a complete, portable Owner workbook from the supplied source files.

    Sources are reparsed even when the caller supplies nodes: exported immutable
    cells must always describe those exact original YAML bytes.
    """
    batch = _source(files)
    nodes = batch.nodes
    tables = _tables(nodes, include_column_tags=VERSION not in (LEGACY_VERSION, PREVIOUS_VERSION))
    output = BytesIO()
    with xlsxwriter.Workbook(output, {"in_memory": True, "strings_to_formulas": False,
                                     "strings_to_urls": False}) as book:
        book.set_properties({"title": "Revisão de metadados pelo Owner",
                             "comments": "Processado integralmente em memória."})
        styles = _styles(book)
        _guide(book, styles, nodes, rules.validate_nodes(nodes))
        list_ws = book.add_worksheet(LIST_SHEET)
        list_column = 0
        list_ranges = {}
        catalog = rules.get_catalog()
        for name, table in tables.items():
            ws = book.add_worksheet(name)
            ws.hide_gridlines(2)
            ws.set_tab_color("#D3A827" if name == "Colunas" else "#407FAE")
            width = len(table["headers"])
            ws.merge_range(0, 0, 0, min(width - 1, 6), name + " — revisão pelo Owner", styles["title"])
            ws.set_row(0, 32)
            ws.merge_range(1, 0, 1, min(width - 1, 6),
                "Amarelo: escreva só as alterações. Azul: comentário opcional. Branco na proposta mantém o original; [REMOVER] exclui um valor existente.", styles["hint"])
            ws.set_row(1, 43)
            for col, (header, col_width) in enumerate(zip(table["headers"], table["widths"])):
                ws.set_column(col, col, col_width)
            ws.set_column(table["hidden_start"], width - 1, 20, None, {"hidden": True})
            ws.write_row(HEADER_ROW, 0, table["headers"], styles["header"])
            ws.set_row(HEADER_ROW, 62 if name == "Colunas" else 43)
            if name == "Colunas":
                # The complete table path remains visible; hide its duplicated
                # catalog/schema columns to leave more space for owner input.
                ws.set_column(0, 1, 22, None, {"hidden": True, "level": 1})
                for index, tag in enumerate(table["column_tags"]):
                    for col in (7 + 2 * index, 8 + 2 * index):
                        ws.write_comment(HEADER_ROW, col, "Tag técnica: " + tag + "\n" +
                                         catalog["Column"].get(tag, {}).get("description", "Tag existente sem regra documentada."))
            for row, record in enumerate(table["rows"], HEADER_ROW + 1):
                tallest = 1
                for col, value in enumerate(record["values"]):
                    editable = record["editable"].get(col)
                    is_absent_tag = (name == "Colunas" and col >= 8 and col < table["hidden_start"] - 1
                                     and (col - 8) % 2 == 0 and not editable)
                    fmt = styles["note" if editable and editable[0] == "note" else "input"] if editable else styles["inactive" if is_absent_tag else "current"]
                    ws.write_string(row, col, value, fmt)
                    if col < table["hidden_start"] and value:
                        approx_width = table["widths"][col] - 2
                        tallest = max(tallest, sum(max(1, math.ceil(len(line) / approx_width)) for line in value.split("\n")))
                    if editable and editable[0] == "tags":
                        rule = catalog.get(record["node"]["level"], {}).get(editable[1], {})
                        allowed = rule.get("values")
                        if allowed:
                            choices = tuple((["Sim", "Não"] if allowed == rules.BOOL else allowed) + [DELETE_TOKEN])
                            if choices not in list_ranges:
                                list_ws.write_column(0, list_column, choices)
                                letter = xl_col_to_name(list_column)
                                list_ranges[choices] = f"='{LIST_SHEET}'!${letter}$1:${letter}${len(choices)}"
                                list_column += 1
                            ws.data_validation(row, col, row, col, {"validate": "list",
                                "source": list_ranges[choices], "ignore_blank": True,
                                "error_title": "Valor inválido", "error_message": "Selecione um valor da lista ou deixe vazio para manter.",
                                "input_title": "Alteração opcional", "input_message": "Vazio mantém; [REMOVER] exclui.",
                                "error_type": "stop"})
                ws.set_row(row, min(300, max(42, 15 * tallest + 8)))
            # Excel can filter protected sheets. Sorting requires full-row moves;
            # importer verifies each row independently and accepts reordered rows.
            if table["rows"]:
                ws.autofilter(HEADER_ROW, 0, HEADER_ROW + len(table["rows"]), width - 1)
            ws.freeze_panes(HEADER_ROW + 1, 4 if name == "Colunas" else 2)
            ws.set_landscape()
            ws.fit_to_pages(1, 0)
            ws.repeat_rows(HEADER_ROW)
            ws.protect("", {"autofilter": True, "sort": True, "select_locked_cells": True,
                             "select_unlocked_cells": True})
        list_ws.very_hidden()
        list_ws.protect()
        _write_payload(book, files)
    result = output.getvalue()
    if len(result) > MAX_XLSX_BYTES:
        raise ValueError("O Excel excede 30 MB; divida o lote em arquivos menores.")
    return result


def _formula_error(archive, worksheet, raw):
    """Locate actual cell formulas without evaluating or echoing their contents."""
    label = worksheet
    try:
        workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {rel.get("Id"): rel.get("Target", "") for rel in relationships}
        for sheet in workbook.iter():
            if sheet.tag.rsplit("}", 1)[-1] != "sheet":
                continue
            rid = sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            target = targets.get(rid, "")
            path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
            if path.lower() == worksheet:
                label = sheet.get("name", worksheet)
                break
    except Exception:
        # The reader will diagnose a corrupt workbook separately. Never reveal
        # untrusted XML fragments as part of a diagnostic.
        pass
    for _event, cell in ElementTree.iterparse(BytesIO(raw), events=("end",)):
        if cell.tag.rsplit("}", 1)[-1] != "c":
            continue
        if any(child.tag.rsplit("}", 1)[-1] == "f" for child in cell):
            coordinate = cell.get("r", "não identificada")
            if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", coordinate):
                coordinate = "não identificada"
            raise ValueError(
                f'Aba "{label}", célula {coordinate}: há uma fórmula, e não um valor preenchido. '
                "No Excel, copie essa célula e use Colar especial > Valores; depois salve e envie novamente. "
                "Se for um texto que começa com =, digite um apóstrofo antes do sinal."
            )
        cell.clear()


def _preflight(content):
    """Bound the ZIP/XML before openpyxl builds objects, without extracting files."""
    if not isinstance(content, bytes) or len(content) > MAX_XLSX_BYTES:
        raise ValueError("Envie um Excel .xlsx de até 30 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            names = [info.filename for info in infos]
            if len(infos) > 1500 or len(names) != len(set(names)):
                raise ValueError("Excel contém entradas duplicadas ou em excesso.")
            if sum(info.file_size for info in infos) > MAX_XML_BYTES:
                raise ValueError("Excel descompactado excede o limite de segurança.")
            count_cells = 0
            for info in infos:
                name = info.filename.lower()
                if (info.flag_bits & 1 or name.startswith("/") or "\\" in name or
                        any(p in (".", "..") for p in name.split("/"))):
                    raise ValueError("Estrutura ZIP insegura no arquivo Excel.")
                if ("vbaproject" in name or "externallink" in name or "/embeddings/" in name
                        or name.endswith((".bin", ".vba"))):
                    raise ValueError("Macros, objetos incorporados e vínculos externos não são aceitos.")
                if info.file_size > MAX_XML_BYTES or (info.file_size > 1024 * 1024 and
                        info.file_size / max(info.compress_size, 1) > 500):
                    raise ValueError("Excel excede a taxa de compactação permitida.")
                if name.endswith((".xml", ".rels")):
                    raw = archive.read(info)
                    # OOXML is UTF-8. Rejecting NUL also closes UTF-16 entity bypasses.
                    if b"\x00" in raw or re.search(br"<!\s*(?:DOCTYPE|ENTITY)", raw, re.I):
                        raise ValueError("Declarações XML externas não são aceitas.")
                    if "/worksheets/" in name and re.search(br"<(?:[A-Za-z_][\w.-]*:)?f(?:\s|/?>)", raw):
                        _formula_error(archive, name, raw)
                    if name.endswith(".rels") and re.search(br"TargetMode\s*=\s*['\"]External['\"]", raw, re.I):
                        raise ValueError("Vínculos externos não são aceitos.")
                    if "/worksheets/" in name:
                        count_cells += len(re.findall(br"<(?:[A-Za-z_][\w.-]*:)?c(?:\s|>)", raw))
                        if count_cells > MAX_CELLS + 10000:
                            raise ValueError("Excel contém células em excesso.")
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                raise ValueError("Arquivo Excel inválido.")
    except (BadZipFile, RuntimeError, OSError, EOFError, NotImplementedError, zlib.error):
        raise ValueError("Arquivo Excel inválido ou corrompido.") from None


def _payload(book):
    if PAYLOAD_SHEET not in book.sheetnames:
        raise ValueError("A planilha não contém os YAMLs de origem. Use o Excel gerado por este aplicativo.")
    ws = book[PAYLOAD_SHEET]
    try:
        if ws.max_column != 1 or ws.max_row > 5000:
            raise ValueError
        rows = ws.iter_rows(max_col=1, values_only=True)
        manifest = json.loads(next(rows)[0])
        length, count = manifest["length"], manifest["chunks"]
        if (manifest["version"] not in SUPPORTED_VERSIONS or manifest.get("encoding") != "zlib+base64"
                or type(length) is not int or not 1 <= length <= MAX_PAYLOAD_BYTES
                or type(count) is not int or not 1 <= count <= 4999
                or ws.max_row != count + 1):
            raise ValueError
        chunks = [row[0] for row in rows]
        if any(not isinstance(chunk, str) or not 0 < len(chunk) <= CHUNK_SIZE for chunk in chunks):
            raise ValueError
        if sum(map(len, chunks)) > MAX_PAYLOAD_BYTES * 4 // 3 + 1024:
            raise ValueError
        packed = base64.b64decode("".join(chunks), validate=True)
        decoder = zlib.decompressobj()
        raw = decoder.decompress(packed, length + 1)
        if (len(raw) != length or not decoder.eof or decoder.unconsumed_tail
                or decoder.unused_data or sha256(raw).hexdigest() != manifest["sha256"]):
            raise ValueError
        payload = json.loads(raw)
        if payload["version"] != manifest["version"] or not isinstance(payload["files"], list):
            raise ValueError
        files = payload["files"]
        if any(not isinstance(f, dict) or set(f) != {"name", "content"} or
               not isinstance(f["name"], str) or not isinstance(f["content"], str) for f in files):
            raise ValueError
        return files
    except (ValueError, KeyError, TypeError, binascii.Error, zlib.error, UnicodeError, RecursionError, StopIteration):
        raise ValueError("Os metadados de origem da planilha estão ausentes, alterados ou corrompidos.") from None


def read_owner_workbook(content: bytes) -> dict:
    """Return only authenticated-by-source editable cells, rejecting all on errors.

    The source hash is an integrity check, not an author signature. A new source
    file may be supplied by a new workbook; metadata governance remains offline.
    """
    _preflight(content)
    try:
        book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False,
                                     keep_links=False)
    except Exception:
        raise ValueError("Arquivo Excel inválido ou corrompido.") from None
    try:
        required = set(VISIBLE_SHEETS + [PAYLOAD_SHEET, LIST_SHEET])
        if set(book.sheetnames) != required:
            raise ValueError("As abas da planilha foram removidas, renomeadas ou acrescentadas. Use o modelo completo.")
        for sheet in book:
            if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS:
                raise ValueError("A planilha excede o limite de linhas ou colunas.")
        files = _payload(book)
        source_version = json.loads(book[PAYLOAD_SHEET].cell(1, 1).value)["version"]
        batch = _source(files)
        tables = _tables(batch.nodes, include_column_tags=source_version not in (LEGACY_VERSION, PREVIOUS_VERSION))
        edits, notes, warnings = [], [], []
        proposals = {}
        nodes_by_id = {node["id"]: node for node in batch.nodes}
        catalog = rules.get_catalog()
        # Original validation is advisory; it must never block a blank proposal.
        warnings.extend({**issue, "source": "original"} for issue in rules.validate_nodes(batch.nodes))
        for name, table in tables.items():
            ws = book[name]
            headers = table["headers"]
            width = len(headers)
            actual_headers = [_text(value) for value in next(ws.iter_rows(
                min_row=HEADER_ROW + 1, max_row=HEADER_ROW + 1, max_col=width, values_only=True))]
            if actual_headers != headers:
                raise ValueError(f"{name}: os cabeçalhos foram alterados. Use o modelo exportado pelo aplicativo.")
            expected = {record["key"]: record for record in table["rows"]}
            id_col = headers.index("ID interno")
            tag_col = headers.index("Tag técnica") if name == "Tags" else None
            seen = set()
            for cells in ws.iter_rows(min_row=HEADER_ROW + 2, max_col=max(width, ws.max_column)):
                if not any(cell.value is not None for cell in cells):
                    continue
                values = [_text(cell.value) for cell in cells]
                key = (values[id_col], values[tag_col]) if tag_col is not None else values[id_col]
                location = f"{name}, linha {cells[0].row}"
                if key in seen:
                    raise ValueError(location + ": registro duplicado; não copie ou insira linhas.")
                seen.add(key)
                record = expected.get(key)
                if record is None:
                    raise ValueError(location + ": registro desconhecido ou identificador alterado.")
                if any(values[width:]):
                    raise ValueError(location + ": há conteúdo em colunas não previstas no modelo.")
                for col, old in enumerate(record["values"]):
                    if col not in record["editable"] and values[col] != old:
                        # v1 was issued with both the short table label and the
                        # qualified path. The stable ID, source hash and every
                        # other immutable cell still have to match exactly.
                        if (source_version == LEGACY_VERSION and name == "Colunas"
                                and headers[col] == "Tabela / visão"
                                and values[col] == record["node"]["path"].rsplit(".", 1)[0].rsplit(".", 1)[-1]):
                            continue
                        raise ValueError(location + f': campo protegido "{headers[col]}" ({cells[col].coordinate}) alterado. '
                                         "Edite somente os campos amarelos e azuis.")
                node = record["node"]
                for col, (field, tag) in record["editable"].items():
                    cell, value = cells[col], values[col]
                    if cell.data_type == "f":
                        raise ValueError(location + ": fórmulas não são aceitas.")
                    if not value.strip():
                        continue
                    if not isinstance(cell.value, str):
                        raise ValueError(location + ": a proposta deve ser texto; use true ou false nas listas.")
                    if len(value) > (10_000 if field == "tags" else 32_000):
                        raise ValueError(location + ": o texto proposto excede o tamanho permitido.")
                    if field == "note":
                        notes.append({"id": node["id"], "path": node["path"], "sheet": name,
                                      "row": cell.row, "tag": tag, "comment": value})
                        continue
                    if value == DELETE_TOKEN:
                        exists = tag in node["tags"] if field == "tags" else node.get(field) is not None
                        if not exists:
                            raise ValueError(location + ": o campo já está ausente; não há valor para remover.")
                        value = None
                    if field == "tags" and value is not None:
                        rule = catalog.get(node["level"], {}).get(tag, {})
                        if rule.get("values") == rules.BOOL:
                            value = {"Sim": "true", "Não": "false"}.get(value, value)
                        tag_issues = rules.validate_tag(node["level"], tag, value, catalog=catalog)
                        errors = [issue for issue in tag_issues if issue["severity"] == "error"]
                        if errors:
                            raise ValueError(location + ", " + TAG_LABELS.get(tag, tag) + ": " + errors[0]["message"])
                        warnings.extend({"id": node["id"], "path": node["path"], **issue,
                                         "source": "proposal"} for issue in tag_issues)
                    edit = {"id": node["id"], "field": field, "value": value}
                    if tag is not None:
                        edit["tag"] = tag
                    key = (node["id"], field, tag)
                    cell_location = f"{name}!{cell.coordinate}"
                    previous = proposals.get(key)
                    if previous is not None and previous[0]["value"] != value:
                        raise ValueError(
                            f"Propostas diferentes para a mesma informação em {previous[1]} e {cell_location}. "
                            "Preencha somente uma das abas ou use o mesmo valor nas duas."
                        )
                    proposals.setdefault(key, (edit, cell_location))
            if seen != set(expected):
                raise ValueError(f"{name}: faltam registros da origem; não exclua linhas do modelo.")
        for edit, _location in proposals.values():
            node = nodes_by_id[edit["id"]]
            field, tag, value = edit["field"], edit.get("tag"), edit["value"]
            old = node["tags"].get(tag) if field == "tags" else node.get(field)
            # Compare only after duplicate proposals are checked, including an
            # explicit original value versus a changed value in the other tab.
            if (value is not None and type(old) is type(value) and old == value
                    and not (field == "tags" and tag in node.get("_native_tag_types", {}))):
                continue
            edits.append(edit)
        return {"files": files, "edits": edits, "notes": notes, "warnings": warnings}
    finally:
        book.close()
